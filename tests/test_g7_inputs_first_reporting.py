"""G7 source-backed advisory outcomes; synthetic fixtures only, PG16 optional."""
from copy import deepcopy
from concurrent.futures import ThreadPoolExecutor
from datetime import date, timedelta
from io import BytesIO
import json
import os
from types import SimpleNamespace
from threading import Barrier
from uuid import uuid4

from openpyxl import load_workbook
import pytest
from sqlalchemy import select
from sqlalchemy.orm import Session

from mcias.audit_agents import AGENTS
from mcias.audit_contract import claim_reference_projection
from mcias.audit_service import run_audit
from mcias.claim_batches import stage, resume
from mcias.governance_g6 import open_case, append_event
from mcias.models import AuditLog, Organization, User
from mcias.models_reporting_g7 import ReportSnapshot
from mcias.models_governance_g6 import GovernanceEvent
from mcias.report_exports_g7 import render_pdf, render_xlsx
from mcias.reporting_g7 import ReportBlocked, create_snapshot, fingerprint, freeze_batch
from mcias import reporting_g7_explorer as explorer
from test_g7_reporting import fixture
from test_g5_inputs_first_review import native_review, _claim, _references, _execution


def audit_fixture(state="MATCH_VERIFIED"):
    batch, claim, snapshot, case, event = fixture()
    claim.id = str(uuid4()); claim.extracted_record_id = str(uuid4())
    snapshot.id = str(uuid4()); snapshot.claim_id = claim.id; snapshot.policy_fingerprint = "a" * 64
    case.claim_id = claim.id; case.audit_snapshot_id = snapshot.id
    event.snapshot_id = snapshot.id
    claim.service_date = date(2026, 6, 1)
    batch.ingestion_document_id = str(uuid4())
    claim.source_locator.update(document_id=batch.ingestion_document_id, source_sha256=batch.source_sha256,
        source_scope_approval_id=str(uuid4()), index=1)
    certificates = [{"family": family, "reference_id": str(uuid4()), "version": 1,
        "organization_id": claim.organization_id, "approved": True, "complete": True,
        "lookup_status": "reference_applicable", "gaps": [],
        "effective_from": "2026-01-01", "effective_until": "2026-07-01",
        "coverage_from": "2026-01-01", "coverage_until": "2026-07-01",
        "source_evidence": [{"source_record_id": str(uuid4())}],
        "records": [{"values": {"beneficiary_token": "PRIVATE-ROSTER-TOKEN"}}]}
        for family in ("beneficiary", "contract", "tariff", "policy", "caps", "history")]
    agents = [{"agent": agent, "rule_id": agent, "rule_version": 3 if agent == "duplicates" else 2,
        "audit_state": "MATCH_VERIFIED", "reference_families": ["policy"], "reference": [str(uuid4())]}
        for agent in AGENTS]
    reason = "CHECKS_COMPLETED"
    coverage = {"required": 8, "executed": 8, "blocked": 0}
    source = {"source_record_id": claim.extracted_record_id, "document_id": batch.ingestion_document_id,
        "source_sha256": batch.source_sha256, "source_scope_approval_id": claim.source_locator["source_scope_approval_id"]}
    if state == "VIOLATION_PROVEN":
        agents[2]["audit_state"] = state
        reason = "PROVEN_RULE_BREACH"
    elif state == "UNRESOLVED":
        certificates[0].update(complete=False, gaps=["reference_unapproved"])
        agents[0]["audit_state"] = state
        coverage.update(executed=7, blocked=1)
        reason = "REFERENCE_INCOMPLETE"
    elif state == "AUDIT_UNAVAILABLE":
        source = None
        agents[-1]["audit_state"] = state
        coverage.update(executed=7, blocked=1)
        reason = "SOURCE_UNAVAILABLE"
    snapshot.findings = {"audit_state": state, "audit_state_schema_version": 2,
        "reason_code": reason, "reference_certificates": certificates,
        "coverage": coverage, "source_evidence": source, "agents": agents}
    return batch, claim, snapshot, case, event


def shared_event(batch, claim, snapshot):
    details = {"organization_id": claim.organization_id, "outcome": "success", "audit_snapshot_id": snapshot.id,
        "document_id": claim.source_locator["document_id"], "source_sha256": batch.source_sha256,
        "source_record_id": claim.extracted_record_id, "policy_fingerprint": snapshot.policy_fingerprint,
        "audit_state": snapshot.findings["audit_state"], "audit_state_schema_version": snapshot.findings.get("audit_state_schema_version", 1),
        "reason_code": snapshot.findings.get("reason_code"),
        "references": claim_reference_projection(snapshot.findings["reference_certificates"])}
    return SimpleNamespace(id=str(uuid4()), action="claim_inputs_audited", resource_type="claim_entry",
        resource_id=claim.id, details=details)


@pytest.mark.parametrize("state", ["MATCH_VERIFIED", "VIOLATION_PROVEN", "UNRESOLVED", "AUDIT_UNAVAILABLE"])
def test_four_outcomes_are_separate_from_financial_authority(state):
    batch, claim, snapshot, case, _ = audit_fixture(state)
    before = deepcopy(snapshot.findings)
    dataset, summary = freeze_batch(batch, [claim], [snapshot], [case], [], [shared_event(batch, claim, snapshot)])
    row = dataset["rows"][0]
    assert row["audit_state"] == state and row["audit_state_schema_version"] == 2
    assert row["completed_audit"] is (state in {"MATCH_VERIFIED", "VIOLATION_PROVEN"})
    assert summary["audit"]["completed_audit_count"] == int(row["completed_audit"])
    assert summary["audit"]["incomplete_audit_count"] == int(not row["completed_audit"])
    assert summary["audit"]["outcome_counts"][state] == 1
    assert row["state"] == "held" and row["payable"] is None and row["deduction"] is None
    assert "PRIVATE-ROSTER-TOKEN" not in json.dumps(row["audit_references"])
    assert "records" not in json.dumps(row["audit_references"])
    assert snapshot.findings == before


@pytest.mark.parametrize("fault", ["certificate", "source", "coverage", "rule", "date", "history_window", "rule_id", "rule_version", "negative_rule_version", "batch_document"])
def test_matching_name_alone_never_proves_a_completed_audit(fault):
    batch, claim, snapshot, case, _ = audit_fixture()
    event = shared_event(batch, claim, snapshot)
    if fault == "certificate":
        snapshot.findings.pop("reference_certificates")
    elif fault == "source":
        snapshot.findings["source_evidence"]["source_sha256"] = "0" * 64
    elif fault == "coverage":
        snapshot.findings["coverage"].update(executed=7, blocked=1)
    elif fault == "rule":
        snapshot.findings["agents"][0]["audit_state"] = "UNRESOLVED"
    elif fault == "date":
        snapshot.findings["reference_certificates"][0]["effective_until"] = "2026-06-01"
    elif fault == "history_window":
        snapshot.findings["agents"][-2].update(audit_state="UNRESOLVED", reason="history_lookback_incomplete")
    elif fault == "rule_id":
        snapshot.findings["agents"][0]["rule_id"] = "tariff"
    elif fault == "rule_version":
        snapshot.findings["agents"][0].pop("rule_version")
    elif fault == "negative_rule_version":
        snapshot.findings["agents"][0]["rule_version"] = -1
    elif fault == "batch_document":
        batch.ingestion_document_id = str(uuid4())
    if fault == "certificate":
        with pytest.raises(ReportBlocked, match="outcome/reference mismatch"):
            freeze_batch(batch, [claim], [snapshot], [case], [], [event])
        return
    dataset, summary = freeze_batch(batch, [claim], [snapshot], [case], [], [shared_event(batch, claim, snapshot)])
    assert dataset["rows"][0]["completed_audit"] is False
    assert summary["audit"]["completed_audit_count"] == 0


def test_legacy_read_projection_preserves_history_and_does_not_invent_reason():
    batch, claim, snapshot, case, _ = audit_fixture()
    snapshot.findings = {"audit_state": "VERIFIED_MATCH", "agents": []}
    original = deepcopy(snapshot.findings)
    dataset, summary = freeze_batch(batch, [claim], [snapshot], [case], [])
    row = dataset["rows"][0]
    assert row["audit_state"] == "MATCH_VERIFIED" and row["source_audit_state_schema_version"] == 1
    assert row["audit_reason_code"] is None and row["completed_audit"] is False
    assert summary["audit"]["incomplete_audit_count"] == 1
    assert snapshot.findings == original
    snapshot.findings = {"agents": []}
    _, summary = freeze_batch(batch, [claim], [snapshot], [case], [])
    assert summary["audit"]["unclassified_count"] == 1
    assert summary["audit"]["incomplete_audit_count"] == 1


def test_latest_snapshot_and_its_events_are_selected_without_mutating_old_report():
    batch, claim, snapshot, case, event = audit_fixture()
    old, old_summary = freeze_batch(batch, [claim], [snapshot], [case], [])
    digest = fingerprint(old, old_summary)
    _, _, newer, _, _ = audit_fixture("UNRESOLVED")
    newer.claim_id = claim.id; newer.id = str(uuid4()); newer.created_at = snapshot.created_at + timedelta(days=1)
    newer_case = SimpleNamespace(**{**vars(case), "id": "case2", "audit_snapshot_id": newer.id})
    newer_event = SimpleNamespace(**{**vars(event), "case_id": newer_case.id, "snapshot_id": newer.id, "kind": "hold"})
    new, summary = freeze_batch(batch, [claim], [snapshot, newer], [case, newer_case], [newer_event])
    assert new["rows"][0]["audit_snapshot_id"] == newer.id
    assert new["rows"][0]["governance_case_id"] == newer_case.id
    assert summary["audit"]["outcome_counts"]["UNRESOLVED"] == 1
    assert old["rows"][0]["audit_state"] == "MATCH_VERIFIED"
    assert fingerprint(old, old_summary) == digest


def test_reference_and_shared_audit_organization_mismatch_fail_closed():
    batch, claim, snapshot, case, _ = audit_fixture()
    snapshot.findings["reference_certificates"][0]["organization_id"] = "foreign"
    with pytest.raises(ReportBlocked, match="cross-organization reference"):
        freeze_batch(batch, [claim], [snapshot], [case], [])
    snapshot.findings["reference_certificates"][0]["organization_id"] = claim.organization_id
    event = SimpleNamespace(action="claim_inputs_audited", details={"organization_id": "foreign"})
    with pytest.raises(ReportBlocked, match="cross-organization audit event"):
        freeze_batch(batch, [claim], [snapshot], [case], [], [event])


def test_completed_audit_requires_durable_shared_audit_correlation():
    batch, claim, snapshot, case, _ = audit_fixture()
    dataset, summary = freeze_batch(batch, [claim], [snapshot], [case], [], [])
    assert dataset["rows"][0]["audit_state"] == "MATCH_VERIFIED"
    assert dataset["rows"][0]["completed_audit"] is False
    assert summary["audit"]["completed_audit_count"] == 0


@pytest.mark.parametrize("fault", ["duplicate", "claim", "document", "sha", "state", "schema", "reference", "resource_type", "outcome"])
def test_conflicting_shared_audit_evidence_blocks_the_report(fault):
    batch, claim, snapshot, case, _ = audit_fixture()
    event = shared_event(batch, claim, snapshot)
    events = [event]
    if fault == "duplicate":
        events.append(shared_event(batch, claim, snapshot))
    elif fault == "claim":
        event.resource_id = str(uuid4())
    elif fault == "document":
        event.details["document_id"] = str(uuid4())
    elif fault == "sha":
        event.details["source_sha256"] = "0" * 64
    elif fault == "state":
        event.details.update(audit_state="VIOLATION_PROVEN", reason_code="PROVEN_RULE_BREACH")
    elif fault == "schema":
        event.details.update(audit_state="VERIFIED_MATCH", audit_state_schema_version=1)
    elif fault == "resource_type":
        event.resource_type = "other_resource"
    elif fault == "outcome":
        event.details["outcome"] = "denied"
    else:
        event.details["references"][0]["version"] = 2
    with pytest.raises(ReportBlocked, match="shared audit"):
        freeze_batch(batch, [claim], [snapshot], [case], [], events)


@pytest.mark.parametrize("legacy", [False, True])
def test_rule_evidence_drops_phi_marker_and_retains_only_canonical_ids(legacy):
    batch, claim, snapshot, case, _ = audit_fixture()
    original_id = snapshot.findings["agents"][0]["reference"][0]
    snapshot.findings["agents"][0]["reference"] += ["PHI-HOLDOUT-PRIVATE-TEXT"]
    if legacy:
        snapshot.findings["audit_state"] = "VERIFIED_MATCH"
        snapshot.findings.pop("audit_state_schema_version")
        for check in snapshot.findings["agents"]:
            check["audit_state"] = "VERIFIED_MATCH"
    dataset, _ = freeze_batch(batch, [claim], [snapshot], [case], [])
    row = dataset["rows"][0]
    assert row["audit_rules"][0]["evidence_record_ids"] == [original_id]
    assert "PHI-HOLDOUT-PRIVATE-TEXT" not in json.dumps(row)
    assert row["completed_audit"] is False


@pytest.mark.parametrize("field", ["gaps", "evidence_ids"])
def test_reference_projection_rejects_phi_without_echoing_it(field):
    batch, claim, snapshot, case, _ = audit_fixture()
    certificate = snapshot.findings["reference_certificates"][0]
    if field == "gaps":
        certificate["gaps"] = ["PHI-HOLDOUT-PRIVATE-TEXT"]
    else:
        certificate["source_evidence"] = [{"source_record_id": "PHI-HOLDOUT-PRIVATE-TEXT"}]
    with pytest.raises(ReportBlocked) as blocked:
        freeze_batch(batch, [claim], [snapshot], [case], [])
    assert "PHI-HOLDOUT-PRIVATE-TEXT" not in str(blocked.value)


def test_redacted_explorer_filters_audit_outcomes_without_roster_or_originals(monkeypatch):
    batch, claim, snapshot, case, _ = audit_fixture()
    dataset, summary = freeze_batch(batch, [claim], [snapshot], [case], [], [shared_event(batch, claim, snapshot)])
    report = SimpleNamespace(id="report1", version=1, organization_id=claim.organization_id,
        batch_id=batch.id, dataset=dataset, summary=summary, fingerprint=fingerprint(dataset, summary))
    db = SimpleNamespace(get=lambda *_: report)
    monkeypatch.setattr(explorer, "authorized", lambda *args: SimpleNamespace(organization_id=claim.organization_id))
    monkeypatch.setattr(explorer, "require_organization", lambda *args: True)
    result = explorer.entries(report.id, state=None, currency=None, limit=10, offset=0,
        db=db, audit_state="MATCH_VERIFIED")
    row = result["items"][0]
    assert row["audit_state"] == "MATCH_VERIFIED" and row["source_position"]["row"] == 7
    assert row["audit_classification"] == "NO_DETECTED_FINDINGS"
    assert row["financial_decision_status"] == "PENDING_HUMAN_REVIEW"
    assert row["audit_references"][0]["effective_from"] == "2026-01-01"
    assert row["audit_references"][0]["coverage_until"] == "2026-07-01"
    assert all(secret not in json.dumps(result) for secret in ("PRIVATE-ROSTER-TOKEN", "token-one", "synthetic-test-only", "source_values", "source_locator"))
    assert explorer.entries(report.id, state=None, currency=None, limit=10, offset=0,
        db=db, audit_state="UNRESOLVED")["total"] == 0
    wb = load_workbook(BytesIO(render_xlsx(report)))
    assert wb.sheetnames == ["Executive Summary"]
    cells = [cell.value for row in wb.active for cell in row]
    assert "MATCH_VERIFIED" in cells and "Completed audits" in cells
    assert "PRIVATE-ROSTER-TOKEN" not in json.dumps(cells)
    assert render_pdf(report).startswith(b"%PDF")


def rejection_events(case, snapshot):
    """The persisted shape of the existing G6 independent-human decision flow."""
    case.id = str(uuid4()); case.opened_by = str(uuid4())
    proposal = SimpleNamespace(id=str(uuid4()), organization_id=case.organization_id,
        case_id=case.id, snapshot_id=snapshot.id, version=1, kind="propose", actor_id=case.opened_by,
        payload={"reason_code": "synthetic_proposal"}, created_at=snapshot.created_at)
    rejection = SimpleNamespace(id=str(uuid4()), organization_id=case.organization_id,
        case_id=case.id, snapshot_id=snapshot.id, version=2, kind="reject", actor_id=str(uuid4()),
        payload={"reason_code": "synthetic_reviewed_rejection"}, created_at=snapshot.created_at + timedelta(minutes=1))
    return [proposal, rejection]


def test_dashboard_four_kpis_five_exclusive_categories_and_separate_currencies():
    fixtures = [audit_fixture(state) for state in
        ("MATCH_VERIFIED", "VIOLATION_PROVEN", "UNRESOLVED", "AUDIT_UNAVAILABLE", "VIOLATION_PROVEN")]
    batch = fixtures[0][0]; batch.imported_count = batch.record_count = 5
    claims, snapshots, cases, events, audits = [], [], [], [], []
    for index, (_, claim, snapshot, case, _) in enumerate(fixtures):
        claim.currency = "USD" if index in (0, 1, 4) else "YER"
        claim.batch_id = batch.id
        claim.source_locator["document_id"] = batch.ingestion_document_id
        if snapshot.findings["source_evidence"]:
            snapshot.findings["source_evidence"]["document_id"] = batch.ingestion_document_id
        case.id = str(uuid4())
        if index == 1:
            snapshot.findings["agents"][4]["audit_state"] = "VIOLATION_PROVEN"
        if index == 4:
            events.extend(rejection_events(case, snapshot))
        claims.append(claim); snapshots.append(snapshot); cases.append(case)
        audits.append(shared_event(batch, claim, snapshot))
    dataset, summary = freeze_batch(batch, claims, snapshots, cases, events, audits)
    dashboard = summary["dashboard"]
    assert len(dashboard["kpis"]) == 4 and len(dashboard["classification_counts"]) == 5
    assert set(dashboard["classification_counts"].values()) == {1}
    assert sum(dashboard["classification_counts"].values()) == summary["total_claims"] == 5
    kpis = dashboard["kpis"]
    assert kpis["total_claimed"]["by_currency"] == {
        "USD": {"status": "KNOWN", "amount": "600", "missing_amount_count": 0},
        "YER": {"status": "KNOWN", "amount": "400", "missing_amount_count": 0}}
    assert kpis["audited_claims"] == {"count": 3, "by_currency": {"USD": 3, "YER": 0}}
    assert kpis["review_required"] == {"claim_count": 2, "finding_count": 3,
        "by_currency": {"USD": {"claim_count": 1, "finding_count": 2}, "YER": {"claim_count": 1, "finding_count": 1}}}
    assert kpis["potential_exposure"]["by_currency"] == {
        currency: {"status": "UNQUANTIFIED", "amount": None, "unquantified_claim_count": 2}
        for currency in ("USD", "YER")}
    assert dashboard["recommended_rejection_pending"] == {"available": False, "claim_count": 0,
        "reason_code": "APPROVED_REJECTION_POLICY_UNAVAILABLE"}
    assert dashboard["human_approved_rejected"] == {"claim_count": 1,
        "claimed_by_currency": {"USD": "200", "YER": "0"}}
    rejected = next(row for row in dataset["rows"] if row["audit_classification"] == "HUMAN_APPROVED_REJECTION")
    assert rejected["decision_event_id"] == events[-1].id
    assert rejected["financial_decision_status"] == "REJECTED"
    assert all(row["state"] == "held" and row["payable"] is None for row in dataset["rows"])


def test_dashboard_matching_advisory_has_no_open_finding_or_financial_approval():
    batch, claim, snapshot, case, _ = audit_fixture()
    snapshot.findings["human_review_required"] = True
    dataset, summary = freeze_batch(batch, [claim], [snapshot], [case], [], [shared_event(batch, claim, snapshot)])
    row = dataset["rows"][0]
    assert row["audit_classification"] == "NO_DETECTED_FINDINGS" and row["review_finding_count"] == 0
    assert row["financial_decision_status"] == "PENDING_HUMAN_REVIEW" and row["payable"] is None
    assert summary["dashboard"]["kpis"]["review_required"]["claim_count"] == 0
    assert summary["dashboard"]["kpis"]["potential_exposure"]["by_currency"]["YER"] == {
        "status": "QUANTIFIED", "amount": "0", "unquantified_claim_count": 0}


def test_dashboard_incomplete_human_rejection_preserves_evidence_gap_and_financial_axis():
    batch, claim, snapshot, case, _ = audit_fixture("UNRESOLVED")
    events = rejection_events(case, snapshot)
    dataset, summary = freeze_batch(batch, [claim], [snapshot], [case], events, [shared_event(batch, claim, snapshot)])
    row = dataset["rows"][0]
    assert row["audit_classification"] == "MISSING_EVIDENCE"
    assert row["financial_decision_status"] == "REJECTED" and row["decision_event_id"] == events[-1].id
    assert row["audit_state"] == "UNRESOLVED" and row["completed_audit"] is False
    assert summary["dashboard"]["human_approved_rejected"]["claim_count"] == 0
    assert summary["dashboard"]["kpis"]["review_required"]["claim_count"] == 1
    assert row["payable"] is None and row["deduction"] is None


@pytest.mark.parametrize("fault", ["same_actor", "no_previous", "wrong_previous", "missing_actor", "missing_time", "missing_reason"])
def test_dashboard_human_rejection_requires_a_qualified_recorded_g6_event(fault):
    batch, claim, snapshot, case, _ = audit_fixture("VIOLATION_PROVEN")
    events = rejection_events(case, snapshot)
    if fault == "same_actor":
        events[-1].actor_id = case.opened_by
    elif fault == "no_previous":
        events = events[-1:]
    elif fault == "wrong_previous":
        events[0].kind = "object"
    elif fault == "missing_actor":
        events[-1].actor_id = None
    elif fault == "missing_time":
        events[-1].created_at = None
    else:
        events[-1].payload = {}
    dataset, summary = freeze_batch(batch, [claim], [snapshot], [case], events, [shared_event(batch, claim, snapshot)])
    row = dataset["rows"][0]
    assert row["audit_classification"] == "HUMAN_REVIEW_REQUIRED"
    assert row["financial_decision_status"] == "PENDING_HUMAN_REVIEW" and row["decision_event_id"] is None
    assert summary["dashboard"]["human_approved_rejected"]["claim_count"] == 0


def test_dashboard_objection_reopens_the_latest_human_decision_without_erasing_history():
    batch, claim, snapshot, case, _ = audit_fixture("VIOLATION_PROVEN")
    events = rejection_events(case, snapshot)
    before, summary = freeze_batch(batch, [claim], [snapshot], [case], events, [shared_event(batch, claim, snapshot)])
    digest = fingerprint(before, summary)
    objection = SimpleNamespace(**{**vars(events[-1]), "id": str(uuid4()), "version": 3,
        "kind": "object", "actor_id": case.opened_by, "created_at": snapshot.created_at + timedelta(minutes=2)})
    after, _ = freeze_batch(batch, [claim], [snapshot], [case], [*events, objection], [shared_event(batch, claim, snapshot)])
    assert before["rows"][0]["audit_classification"] == "HUMAN_APPROVED_REJECTION"
    assert after["rows"][0]["audit_classification"] == "HUMAN_REVIEW_REQUIRED"
    assert after["rows"][0]["financial_decision_status"] == "PENDING_HUMAN_REVIEW"
    assert after["rows"][0]["decision_event_id"] is None
    assert fingerprint(before, summary) == digest


def test_dashboard_latest_snapshot_does_not_inherit_an_older_rejection():
    batch, claim, snapshot, case, _ = audit_fixture("VIOLATION_PROVEN")
    events = rejection_events(case, snapshot)
    newer = deepcopy(snapshot); newer.id = str(uuid4()); newer.created_at += timedelta(days=1)
    newer.findings.update(audit_state="MATCH_VERIFIED", reason_code="CHECKS_COMPLETED")
    for check in newer.findings["agents"]:
        check["audit_state"] = "MATCH_VERIFIED"
    dataset, _ = freeze_batch(batch, [claim], [snapshot, newer], [case], events,
        [shared_event(batch, claim, snapshot), shared_event(batch, claim, newer)])
    row = dataset["rows"][0]
    assert row["audit_snapshot_id"] == newer.id and row["audit_classification"] == "NO_DETECTED_FINDINGS"
    assert row["financial_decision_status"] == "NOT_DECIDED" and row["decision_event_id"] is None


def test_dashboard_does_not_infer_exposure_or_rejection_policy_from_financial_differences():
    from test_g7_reporting import frozen
    report = frozen()
    row = report.dataset["rows"][0]
    assert row["state"] == "approved" and row["deduction"] == "136.00" and row["payable"] == "64.00"
    assert row["audit_classification"] == "UNABLE_TO_AUDIT" and row["financial_decision_status"] == "APPROVED"
    exposure = report.summary["dashboard"]["kpis"]["potential_exposure"]["by_currency"]["YER"]
    assert exposure["status"] == "UNQUANTIFIED" and exposure["amount"] is None
    batch, claim, snapshot, case, _ = audit_fixture("VIOLATION_PROVEN")
    snapshot.findings.update(affected_amount="200", immediate_rejection_recommended=True,
        rejection_policy={"approved": True, "reason_code": "TARIFF_REJECTION_POLICY_BREACH"})
    _, summary = freeze_batch(batch, [claim], [snapshot], [case], [], [shared_event(batch, claim, snapshot)])
    assert summary["dashboard"]["recommended_rejection_pending"]["available"] is False
    assert summary["dashboard"]["recommended_rejection_pending"]["claim_count"] == 0
    assert summary["dashboard"]["kpis"]["potential_exposure"]["by_currency"]["YER"]["amount"] is None


def test_dashboard_cannot_classify_a_foreign_organization_rejection():
    batch, claim, snapshot, case, _ = audit_fixture("VIOLATION_PROVEN")
    events = rejection_events(case, snapshot); events[-1].organization_id = "foreign"
    with pytest.raises(ReportBlocked, match="cross-organization"):
        freeze_batch(batch, [claim], [snapshot], [case], events, [shared_event(batch, claim, snapshot)])


@pytest.mark.parametrize("scenario,expected", [("match", "MATCH_VERIFIED"), ("tariff", "VIOLATION_PROVEN"),
    ("absent_incomplete", "UNRESOLVED"), ("unreadable_source", "AUDIT_UNAVAILABLE")])
def test_native_services_reach_g7_with_durable_safe_lineage(native_review, scenario, expected):
    db, actor, store, engine = native_review
    claim, document, batch = _claim(db, actor, store, price="120" if scenario == "tariff" else "100",
        copay="24" if scenario == "tariff" else "20")
    references = _references(db, actor, store, scenario)
    _execution(db, actor, store, claim, scenario)
    if scenario == "unreadable_source":
        document.status = "failed"; db.commit()
    snapshot = run_audit(db, actor, claim.id, review_only=True)
    case = open_case(db, actor, snapshot.id)
    hold = append_event(db, actor, case.id, "hold", {"reason_code": "advisory_human_review"})
    if engine.dialect.name == "postgresql" and scenario == "match":
        barrier = Barrier(2)
        def concurrent_freeze():
            with Session(engine, expire_on_commit=False) as parallel:
                reviewer = parallel.get(User, actor.id)
                barrier.wait(timeout=30)
                frozen = create_snapshot(parallel, reviewer, batch.id)
                return frozen.id, frozen.version, frozen.fingerprint
        with ThreadPoolExecutor(max_workers=2) as workers:
            futures = [workers.submit(concurrent_freeze) for _ in range(2)]
            results = [future.result(timeout=120) for future in futures]
        assert results[0] == results[1]
        report = db.get(ReportSnapshot, results[0][0])
    else:
        report = create_snapshot(db, actor, batch.id)
    row = report.dataset["rows"][0]
    assert row["audit_state"] == expected
    assert row["completed_audit"] is (expected in {"MATCH_VERIFIED", "VIOLATION_PROVEN"})
    assert (row["state"], row["payable"], row["deduction"]) == ("held", None, None)
    assert row["source_sha256"] == document.sha256 and row["source_record_id"] == claim.extracted_record_id
    assert row["source_document_id"] == document.id and row["source_position"]["index"] == 1
    assert row["source_position"]["row"] == 2
    assert row["governance_case_id"] == case.id and row["governance_event_id"] == hold.id
    assert row["audit_snapshot_id"] == snapshot.id and len(row["audit_event_ids"]) == 1
    assert len(row["audit_references"]) == 6 and len(row["audit_rules"]) == 8
    assert all("records" not in reference for reference in row["audit_references"])
    if scenario == "match":
        assert references[-1].version == 2
        assert next(item for item in row["audit_references"] if item["family"] == "tariff")["version"] == 1
    assert create_snapshot(db, actor, batch.id).id == report.id
    if scenario != "unreadable_source":
        assert stage(db, actor, document.id).id == batch.id
        assert resume(db, actor, batch.id)["status"] == "imported"
    assert create_snapshot(db, actor, batch.id).id == report.id
    assert len(db.scalars(select(ReportSnapshot)).all()) == 1
    report_events = db.scalars(select(AuditLog).where(AuditLog.action == "g7_report_frozen")).all()
    assert len(report_events) == 1 and report_events[0].resource_id == report.id
    assert report_events[0].details == {"organization_id": actor.organization_id, "outcome": "success", "version": 1}
    with Session(engine) as reread:
        durable = reread.get(ReportSnapshot, report.id)
        assert durable.dataset == report.dataset and durable.summary == report.summary
        assert reread.get(AuditLog, row["audit_event_ids"][0]).details["audit_snapshot_id"] == snapshot.id
        assert reread.get(GovernanceEvent, hold.id).snapshot_id == snapshot.id
    wb = load_workbook(BytesIO(render_xlsx(report)))
    cells = [cell.value for worksheet in wb for line in worksheet for cell in line]
    assert expected in cells and "SYN-BEN" not in json.dumps(cells) and "SYN-CLAIM" not in json.dumps(cells)
    assert render_pdf(report).startswith(b"%PDF")
    other = Organization(name="SYN-G7-OTHER"); db.add(other); db.flush()
    outsider = User(organization_id=other.id, username="SYN-G7-" + uuid4().hex, password_hash="synthetic")
    db.add(outsider); db.commit()
    with pytest.raises(ReportBlocked, match="accessible"):
        create_snapshot(db, outsider, batch.id)
    previous_fingerprint = report.fingerprint
    later = append_event(db, actor, case.id, "hold", {"reason_code": "advisory_human_review"})
    newer = create_snapshot(db, actor, batch.id)
    assert newer.version == 2 and newer.dataset["rows"][0]["governance_event_id"] == later.id
    assert report.fingerprint == previous_fingerprint
    assert report.dataset["rows"][0]["governance_event_id"] == hold.id
    output = os.getenv("MCIAS_G7_EVIDENCE_OUTPUT")
    if output:
        with open(output, "a", encoding="utf-8") as evidence:
            evidence.write(json.dumps({"scenario": scenario, "expected": expected, "actual": row["audit_state"],
                "database": engine.dialect.name, "report_id": report.id, "report_version": report.version,
                "report_fingerprint": report.fingerprint, "summary": report.summary,
                "claim_id": claim.id, "snapshot_id": snapshot.id, "audit_event_ids": row["audit_event_ids"],
                "source_record_id": row["source_record_id"], "source_document_id": document.id,
                "source_sha256": document.sha256, "source_position": row["source_position"],
                "governance_case_id": case.id, "governance_event_id": hold.id,
                "references": row["audit_references"], "financial_state": row["state"], "payable": row["payable"]}) + "\n")
