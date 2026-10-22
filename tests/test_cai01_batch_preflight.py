"""Synthetic source-accounting tests, no patient data or policy approvals."""
from types import SimpleNamespace
import pytest
from mcias.claims_batch_preflight import inspect_extracted_records

VALID = dict(claim_reference='SYN-1',beneficiary_token='SYN-B',provider_code='SYN-P',
             service_code='SYN-S',service_date='2026-10-20',claimed_amount='10.0000',currency='YER')

def row(n, *, sheet=None, fields=None, record_id=None):
    return SimpleNamespace(id=record_id or f'SYN-R{n}',
        locator=dict(kind='sheet' if sheet else 'row',index=sheet or n,row=n+1),
        values=fields if fields is not None else dict(VALID, claim_reference=f'SYN-{n}'))

def test_multisheet_requires_attested_scope():
    source=[row(1,sheet=1),row(2,sheet=1),row(3,sheet=2,fields={'amount':'20'}),row(4,sheet=3,fields={'total':'10'})]
    blocked=inspect_extracted_records(source,expected_records=4)
    assert blocked['status']=='BLOCKED'
    assert blocked['counts']['unscoped']==4
    assert 'CLAIM_SHEET_SELECTION_REQUIRED' in blocked['reason_codes']
    result=inspect_extracted_records(source,claim_sheet_indices={1},expected_records=4)
    assert result['status']=='REVIEW_REQUIRED'
    assert result['counts']==dict(source=4,included=2,excluded=2,unscoped=0,invalid=0,ready=2)
    assert len(result['possible_duplicate_record_ids'])==1
    assert result['semantics']=='READ_ONLY_PREFLIGHT_NOT_FINANCIAL_AUDIT'

def test_invalid_claim_row_kept_as_explicit_hold():
    source=[row(1,fields=VALID),row(2,fields={**VALID,'claimed_amount':'NaN'})]
    result=inspect_extracted_records(source,expected_records=2)
    assert result['status']=='BLOCKED'
    assert result['counts']['invalid']==1
    assert result['counts']['ready']==1
    assert result['row_outcomes'][1]['status']=='HOLD_INVALID'

def test_missing_extracted_record_count_blocks():
    result=inspect_extracted_records([row(1)],expected_records=2)
    assert result['status']=='BLOCKED'
    assert 'SOURCE_COUNT_MISMATCH' in result['reason_codes']

def test_selection_must_identify_real_sheet():
    result=inspect_extracted_records([row(1,sheet=1)],claim_sheet_indices={2})
    assert 'UNKNOWN_CLAIM_SHEET' in result['reason_codes']
    assert result['status']=='BLOCKED'

@pytest.mark.parametrize('selection',[[],[0],['1'],[1,1]])
def test_invalid_sheet_selection_rejected(selection):
    with pytest.raises(ValueError):
        inspect_extracted_records([row(1,sheet=1)],claim_sheet_indices=selection)

def test_no_auto_duplicate_violation_or_savings():
    a=row(1,fields=dict(VALID,claim_reference='SYN-A'))
    b=row(2,fields=dict(VALID,claim_reference='SYN-B'))
    result=inspect_extracted_records([a,b],expected_records=2)
    assert result['status']=='REVIEW_REQUIRED'
    assert result['possible_duplicate_record_ids']==[['SYN-R1','SYN-R2']]
    assert all('deduction' not in str(v).lower() for v in result.values())

def test_large_batch_complete_accounting():
    source=[row(i) for i in range(1,2001)]
    result=inspect_extracted_records(source,expected_records=2000)
    assert result['counts']==dict(source=2000,included=2000,excluded=0,unscoped=0,invalid=0,ready=2000)
    assert len(result['row_outcomes'])==2000
    assert result['status']=='REVIEW_REQUIRED'

def test_source_integrity_duplicate_locator_blocked():
    a=row(1)
    b=row(1,record_id='SYN-OTHER')
    result=inspect_extracted_records([a,b],expected_records=2)
    assert result['status']=='BLOCKED'
    assert 'DUPLICATE_SOURCE_ID_OR_LOCATOR' in result['reason_codes']

def test_actual_g2_xlsx_parser_three_sheets_no_double_count():
    from io import BytesIO
    from openpyxl import Workbook
    from mcias.ingestion.parsers.excel_parser import parse
    from mcias.ingestion.provenance import locator_dict
    book = Workbook()
    details = book.active
    details.title = 'Claims Details'
    headers = list(VALID)
    details.append(headers)
    details.append([VALID[k] for k in headers])
    details.append([dict(VALID, claim_reference='SYN-2')[k] for k in headers])
    payment = book.create_sheet('Payment')
    payment.append(['payment_id', 'total'])
    payment.append(['SYN-SETTLEMENT', '20.0000'])
    family = book.create_sheet('Family Summary')
    family.append(['family_id', 'total'])
    family.append(['SYN-FAMILY', '20.0000'])
    data = BytesIO()
    book.save(data)
    extraction = parse(data.getvalue())
    assert extraction.complete and len(extraction.records) == 4
    records = [SimpleNamespace(id=f'SYN-R{i}', locator=locator_dict(rec.locator), values=rec.values)
               for i, rec in enumerate(extraction.records, 1)]
    result = inspect_extracted_records(records, claim_sheet_indices={1}, expected_records=4)
    assert result['status'] == 'REVIEW_REQUIRED'
    assert result['counts'] == dict(source=4, included=2, excluded=2, unscoped=0, invalid=0, ready=2)
    assert set(result['by_sheet']) == {'1', '2', '3'}
    assert sum(x['included'] for x in result['by_sheet'].values()) == 2
