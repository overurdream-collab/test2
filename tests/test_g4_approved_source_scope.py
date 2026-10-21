"""Approved XLSX scope is source-bound, fail-closed and excludes valid PAYMENT rows."""
from io import BytesIO

import pytest
from openpyxl import Workbook
from sqlalchemy import create_engine, select
from sqlalchemy.orm import Session

from mcias.models import Base, Organization, User, SourceFile, Role, Permission, UserRole, RolePermission
from mcias.models_ingestion import IngestionDocument, IngestionJob, ExtractedSourceRecord
from mcias.models_claims import ClaimEntry
from mcias.ingestion.parsers.excel_parser import parse
from mcias.ingestion.manifest import manifest
from mcias.ingestion.provenance import locator_dict
from mcias.claim_batches import approve_source_scope, source_scope, stage, preview, resume, BatchConflict


@pytest.fixture
def source():
    engine = create_engine("sqlite+pysqlite:///:memory:")
    Base.metadata.create_all(engine)
    with Session(engine, expire_on_commit=False) as db:
        org = Organization(name="SYN-source-scope")
        actor = User(username="SYN-source-reviewer", password_hash="fixture")
        role, permission = Role(code="SYN-importer"), Permission(code="claims.import")
        db.add_all([org, actor, role, permission])
        db.flush()
        actor.organization_id = org.id
        db.add_all([UserRole(user_id=actor.id, role_id=role.id),
                    RolePermission(role_id=role.id, permission_id=permission.id)])
        book = Workbook()
        for sheet in (book.active, book.create_sheet("PAYMENT")):
            if sheet is book.active:
                sheet.title = "CLAIMS"
            sheet.append(["claim_reference", "beneficiary_token", "provider_code", "service_code",
                          "service_date", "claimed_amount", "currency"])
            sheet.append([sheet.title + "-SYN", "SYN-B", "SYN-P", "SYN-S", "2026-09-15", "12.50", "USD"])
        output = BytesIO()
        book.save(output)
        book.close()
        extraction = parse(output.getvalue())
        import hashlib
        digest = hashlib.sha256(output.getvalue()).hexdigest()
        original = SourceFile(sha256=digest, original_filename="synthetic.xlsx", storage_key="SYN-only")
        db.add(original)
        db.flush()
        doc = IngestionDocument(organization_id=org.id, source_id=original.id, sha256=digest,
                                filename="synthetic.xlsx", status="extracted")
        db.add(doc)
        db.flush()
        job = IngestionJob(organization_id=org.id, document_id=doc.id, status="extracted",
                           manifest=manifest(extraction))
        db.add(job)
        db.flush()
        db.add_all([ExtractedSourceRecord(job_id=job.id, locator=locator_dict(r.locator),
                                          values=r.values, confidence=100) for r in extraction.records])
        db.commit()
        yield db, actor, doc, job
    engine.dispose()


def test_payment_valid_for_normalization_still_excluded_and_preserved(source):
    db, actor, doc, job = source
    with pytest.raises(BatchConflict, match="approved source scope"):
        stage(db, actor, doc.id)
    approval = approve_source_scope(db, actor, doc.id, {1: "CLAIMS", 2: "PAYMENT"})
    assert approve_source_scope(db, actor, doc.id, {1: "CLAIMS", 2: "PAYMENT"}).id == approval.id
    batch = stage(db, actor, doc.id)
    assert preview(db, actor, batch.id)["excluded_count"] == 1
    result = resume(db, actor, batch.id)
    assert (result["imported_count"], result["record_count"], result["excluded_count"]) == (1, 1, 1)
    assert result["status"] == "imported"
    assert db.scalar(select(ClaimEntry)).claim_reference == "CLAIMS-SYN"
    assert db.query(ExtractedSourceRecord).count() == 2
    assert source_scope(db, actor, doc.id)["excluded_records"] == 1


@pytest.mark.parametrize("classifications", [{1: "CLAIMS", 2: "CLAIMS"},
    {1: "CLAIMS", 2: "UNKNOWN"}, {1: "CLAIMS"}, {1: "CLAIMS", 2: "PAYMENT", 3: "OTHER"}])
def test_nonclaim_relabel_unknown_or_partial_scope_cannot_authorize(source, classifications):
    db, actor, doc, job = source
    with pytest.raises(BatchConflict):
        approve_source_scope(db, actor, doc.id, classifications)
    assert db.query(ClaimEntry).count() == 0


def test_approval_requires_real_actor_permission(source):
    db, actor, doc, job = source
    unprivileged = User(username="SYN-unprivileged", password_hash="fixture", organization_id=actor.organization_id)
    db.add(unprivileged)
    db.commit()
    with pytest.raises(BatchConflict, match="claims.import"):
        approve_source_scope(db, unprivileged, doc.id, {1: "CLAIMS", 2: "PAYMENT"})


@pytest.mark.parametrize("change", ["values", "name", "inventory", "original_sha"])
def test_source_change_invalidates_approval_even_after_completed_import(source, change):
    db, actor, doc, job = source
    approve_source_scope(db, actor, doc.id, {1: "CLAIMS", 2: "PAYMENT"})
    batch = stage(db, actor, doc.id)
    resume(db, actor, batch.id)
    record = db.scalar(select(ExtractedSourceRecord).where(ExtractedSourceRecord.job_id == job.id))
    if change == "values":
        record.values = {**record.values, "claimed_amount": "999.00"}
    elif change == "name":
        record.locator = {**record.locator, "sheet_name": "RENAMED"}
    elif change == "inventory":
        job.manifest = {**job.manifest, "worksheet_inventory": [
            {**s, "name": "RENAMED"} if s["index"] == 1 else s
            for s in job.manifest["worksheet_inventory"]]}
    else:
        db.get(SourceFile, doc.source_id).sha256 = "f" * 64
    db.commit()
    for operation in (lambda: stage(db, actor, doc.id), lambda: preview(db, actor, batch.id),
                      lambda: resume(db, actor, batch.id), lambda: source_scope(db, actor, doc.id)):
        with pytest.raises(BatchConflict):
            operation()
    assert db.query(ClaimEntry).count() == 1


def test_approved_scope_is_immutable_not_client_reclassifiable(source):
    db, actor, doc, job = source
    approve_source_scope(db, actor, doc.id, {1: "CLAIMS", 2: "PAYMENT"})
    with pytest.raises(BatchConflict, match="immutable"):
        approve_source_scope(db, actor, doc.id, {1: "CLAIMS", 2: "OTHER"})
