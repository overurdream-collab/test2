"""G7 T13/T18/T20 golden data and privacy regressions; no real patient data."""
from datetime import datetime, timezone, timedelta
from decimal import Decimal
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import MagicMock
from hashlib import sha256
from fastapi import HTTPException
from openpyxl import load_workbook
import pytest
from mcias.audit_engine import AuditPolicy, evaluate
from mcias.reporting_g7 import (freeze_batch, fingerprint, ReportBlocked,
                                verify_snapshot, beneficiary_statement)
from mcias.report_exports_g7 import render_pdf, render_xlsx
from mcias import reporting_g7_api as api

T = datetime(2026, 10, 10, tzinfo=timezone.utc)

def fixture():
    policy = AuditPolicy(policy_id="synthetic-signed-fixture", currency="YER",
        rounding="HALF_UP", precision=2,
        order=("exclusion", "tariff", "cap", "copay"),
        copay_percent=Decimal("20"), cap=Decimal("80"),
        exclusion=Decimal("10"), tariff=Decimal("100"), approved=True)
    calc = evaluate(Decimal("200"), "YER", policy, eligible=True,
        evidence_complete=True, source_id="record1", contract_id="contract1",
        beneficiary_enrollment_id="member1")
    batch = SimpleNamespace(id="b1", organization_id="org1", status="imported",
        record_count=1, imported_count=1, source_sha256="d"*64)
    claim = SimpleNamespace(id="c1", batch_id="b1", organization_id="org1",
        claimed_amount=Decimal("200"), currency="YER", claim_reference="REF01",
        extracted_record_id="record1", beneficiary_token="token-one",
        source_values={"raw_formula": "=SUM(1,2)", "raw_identifier": "0012300",
                       "diagnosis": "synthetic-test-only"},
        source_locator={"kind":"sheet", "row": 7})
    snap = SimpleNamespace(id="s1", organization_id="org1", claim_id="c1",
        created_at=T, calculation=calc, verification={"status":"verified"})
    case = SimpleNamespace(id="g1", organization_id="org1", claim_id="c1",
        audit_snapshot_id="s1", created_at=T)
    event = SimpleNamespace(id="e2", organization_id="org1", case_id="g1",
        snapshot_id="s1", version=2, kind="approve")
    return batch, claim, snap, case, event

def frozen(*, override_events=None, override_snapshots=None):
    batch, claim, snap, case, event = fixture()
    dataset, summary = freeze_batch(batch, [claim], override_snapshots or [snap],
        [case], [event] if override_events is None else override_events)
    report = SimpleNamespace(id="report1", batch_id="b1", version=1,
        organization_id="org1", dataset=dataset, summary=summary,
        fingerprint=fingerprint(dataset, summary))
    return report

def test_t13_old_snapshot_never_changes_after_new_audit_and_governance():
    before = frozen()
    original_fingerprint = before.fingerprint
    batch, claim, snap, case, event = fixture()
    revised = SimpleNamespace(**{**vars(snap), "id":"s2", "created_at":T+timedelta(days=1),
                                "verification":{"status":"held"}})
    after_dataset, after_summary = freeze_batch(batch, [claim], [snap, revised],
                                                 [case], [event])
    assert after_summary["by_currency"]["YER"]["held_count"] == 1
    assert after_summary["by_currency"]["YER"]["payable"] == "0"
    assert fingerprint(after_dataset, after_summary) != original_fingerprint
    assert before.summary["by_currency"]["YER"]["payable"] == "64.00"
    assert before.fingerprint == original_fingerprint

def test_t18_original_columns_and_excel_formula_remain_literal():
    report = frozen()
    data = render_xlsx(report, include_originals=True)
    wb = load_workbook(BytesIO(data), data_only=False)
    raw = wb["Original Source Values"]
    values = {raw.cell(1, c).value: raw.cell(2, c).value
              for c in range(4, raw.max_column + 1)}
    assert values["raw_identifier"] == "0012300"
    assert values["raw_formula"] == "=SUM(1,2)"
    assert all(raw.cell(2, c).data_type != "f" for c in range(4, raw.max_column + 1))
    assert report.dataset["rows"][0]["source_locator"] == {"kind":"sheet","row":7}

def test_t20_approved_vs_held_payable_and_reconciliation():
    report = frozen()
    row = report.dataset["rows"][0]
    assert (row["claimed"], row["deduction"], row["payable"]) == ("200", "136.00", "64.00")
    assert report.summary["by_currency"]["YER"]["approved_count"] == 1
    assert report.summary["by_currency"]["YER"]["held_claimed"] == "0"
    assert verify_snapshot(report)

def test_t20_objection_is_held_not_approved():
    batch, claim, snap, case, event = fixture()
    objection = SimpleNamespace(**{**vars(event), "id":"e3", "version":3, "kind":"object"})
    dataset, summary = freeze_batch(batch,[claim],[snap],[case],[event,objection])
    assert dataset["rows"][0]["state"]=="held"
    assert dataset["rows"][0]["hold_reason"]=="open_objection"
    assert summary["by_currency"]["YER"]["payable"]=="0"

def test_t20_invalid_calculated_claim_is_held_not_payable():
    batch, claim, snap, case, event = fixture()
    snap.calculation = {**snap.calculation, "payable":"65.00"}
    dataset, summary = freeze_batch(batch,[claim],[snap],[case],[event])
    assert dataset["rows"][0]["state"] == "held"
    assert summary["by_currency"]["YER"]["payable"] == "0"

def test_t20_foreign_tenant_cannot_merge():
    batch, claim, snap, case, event = fixture()
    snap.organization_id = "other"
    with pytest.raises(ReportBlocked, match="cross-organization"):
        freeze_batch(batch,[claim],[snap],[case],[event])

def test_incomplete_batch_fails_closed():
    batch, claim, snap, case, event = fixture()
    batch.status = "review_required"
    with pytest.raises(ReportBlocked, match="not fully imported"):
        freeze_batch(batch,[claim],[snap],[case],[event])

def test_frozen_dataset_tampering_detected():
    report = frozen()
    report.dataset["rows"][0]["payable"] = "64000"
    with pytest.raises(ReportBlocked, match="integrity"):
        verify_snapshot(report)

def test_default_exports_hide_all_source_identifiers():
    report = frozen()
    excel = render_xlsx(report)
    wb = load_workbook(BytesIO(excel))
    assert wb.sheetnames == ["Executive Summary"]
    pdf = render_pdf(report)
    assert pdf.startswith(b"%PDF")
    assert b"synthetic-test-only" not in pdf
    assert b"token-one" not in pdf

def test_beneficiary_statement_is_bound_to_frozen_version():
    report = frozen()
    statement = beneficiary_statement(report, "token-one")
    assert statement["report_version"] == 1
    assert statement["items"][0]["payable"] == "64.00"
    assert beneficiary_statement(report, "different")["items"] == []

def test_t20_original_export_is_denied_without_approval(monkeypatch):
    report = frozen()
    monkeypatch.setattr(api, "_actor", lambda *_: SimpleNamespace(id="a", organization_id="org1"))
    monkeypatch.setattr(api, "_report", lambda *_: report)
    def check(db,actor,organization,permission):
        if permission == "reports.originals.export":
            raise HTTPException(403, "forbidden")
        return True
    monkeypatch.setattr(api,"require_organization",check)
    with pytest.raises(HTTPException) as error:
        api.download_excel("report1", include_originals=True, db=MagicMock())
    assert error.value.status_code == 403

def test_t20_source_fields_require_individual_permissions(monkeypatch):
    report = frozen()
    monkeypatch.setattr(api, "_actor", lambda *_: SimpleNamespace(id="a", organization_id="org1"))
    monkeypatch.setattr(api, "_report", lambda *_: report)
    monkeypatch.setattr(api, "require_organization", lambda *_: True)
    def deny(db,actor,org,field):
        if field == "diagnosis":
            raise HTTPException(403, "forbidden")
    monkeypatch.setattr(api, "require_sensitive_field", deny)
    with pytest.raises(HTTPException) as error:
        api.download_excel("report1", include_originals=True, db=MagicMock())
    assert error.value.status_code == 403

def test_export_pdf_and_excel_verify_integrity_first():
    report = frozen()
    report.summary["total_claims"] = 100
    with pytest.raises(ReportBlocked):
        render_xlsx(report)
    with pytest.raises(ReportBlocked):
        render_pdf(report)

def test_t18_arabic_rtl_workbook_preserves_literal_originals():
    report = frozen()
    excel = render_xlsx(report, include_originals=True, language="ar")
    wb = load_workbook(BytesIO(excel), data_only=False)
    assert wb.sheetnames == ["الملخص التنفيذي", "القيم الأصلية"]
    assert wb["الملخص التنفيذي"].sheet_view.rightToLeft
    assert wb["القيم الأصلية"].sheet_view.rightToLeft
    raw = wb["القيم الأصلية"]
    names = {raw.cell(1, i).value: i for i in range(1, raw.max_column + 1)}
    assert raw.cell(2, names["raw_identifier"]).value == "0012300"
    assert raw.cell(2, names["raw_formula"]).value == "=SUM(1,2)"
    assert raw.cell(2, names["raw_formula"]).data_type == "s"
