"""No DB/ingestion replay: guards and measurement arithmetic for the benchmark."""
from io import BytesIO
import json

import pytest
from openpyxl import load_workbook

from scripts.benchmark_inputs_first import (
    FAMILIES, ratios, reference_rows, synthetic_rows, timing_partition,
    validate_target, workbook_bytes,
)


LOCAL = "postgresql+psycopg://synthetic:synthetic@127.0.0.1:5432/mcias_benchmark_inputs"


@pytest.mark.parametrize("url,allow,sizes,deadline", [
    (LOCAL, "NO", [10], 300),
    (LOCAL.replace("127.0.0.1", "db.example.com"), "YES", [10], 300),
    (LOCAL.replace("mcias_benchmark_inputs", "mcias_production"), "YES", [10], 300),
    ("sqlite://", "YES", [10], 300),
    (LOCAL.replace("postgresql+psycopg", "postgresql"), "YES", [10], 300),
    (LOCAL + "?host=remote", "YES", [10], 300),
    (LOCAL.replace("synthetic:synthetic@", "synthetic@"), "YES", [10], 300),
    (LOCAL, "YES", [1000], 300),
    (LOCAL, "YES", [10, 10], 300),
    (LOCAL, "YES", [25, 10], 300),
    (LOCAL, "YES", [], 300),
    (LOCAL, "YES", [True], 300),
    (LOCAL, "YES", [10], 301),
    (LOCAL, "YES", [10], 14),
])
def test_fail_closed_target_and_resource_bounds(url, allow, sizes, deadline):
    with pytest.raises(ValueError):
        validate_target(url, allow, sizes, deadline)


def test_explicit_bounded_local_target():
    target = validate_target(LOCAL, "YES", [10, 25, 100], 300)
    assert target.database == "mcias_benchmark_inputs"


def test_nested_timing_partition_does_not_double_count():
    result = timing_partition(20.0, 15.0, 9.0)
    assert result["rules_and_source_validation_seconds"] == 6.0
    assert result["persistence_and_orchestration_seconds"] == 5.0
    assert sum(result[key] for key in ("reference_seconds", "rules_and_source_validation_seconds",
        "persistence_and_orchestration_seconds")) == result["g5_total_seconds"]
    with pytest.raises(ValueError):
        timing_partition(1, 2, 0.5)


def test_ratios_have_actual_denominators_and_reject_invalid_counts():
    assert ratios(7, 8) == {"numerator": 7, "denominator": 8, "fraction": 0.875}
    for args in ((0, 0), (9, 8), (-1, 8)):
        with pytest.raises(ValueError):
            ratios(*args)


@pytest.mark.parametrize("size", [10, 25, 100])
def test_known_matching_sources_are_independent_and_payment_is_separate(size):
    rows = synthetic_rows(size)
    assert len({row["beneficiary_token"] for row in rows}) == size
    assert len({row["claim_reference"] for row in rows}) == size
    book = load_workbook(BytesIO(workbook_bytes(rows)), read_only=True)
    assert book.sheetnames == ["CLAIMS", "PAYMENT"]
    assert book["CLAIMS"].max_row == size + 1
    assert book["PAYMENT"].max_row == 2
    assert book["PAYMENT"]["A2"].value == "SYN-PAYMENT"
    book.close()
    references = reference_rows(rows)
    assert tuple(references) == FAMILIES
    assert sum(len(items) for items in references.values()) == 3 * size + 3
    for family in ("beneficiary", "caps", "history"):
        payloads = [json.loads(item["payload"]) for item in references[family]]
        assert {value["beneficiary_token"] for value in payloads} == {row["beneficiary_token"] for row in rows}
    assert all(json.loads(item["payload"])["consumed_amount"] == "0" for item in references["caps"])
    assert all(json.loads(item["payload"])["service_code"] == "SYN-OTHER" for item in references["history"])
