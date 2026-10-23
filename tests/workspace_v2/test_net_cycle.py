"""Complete real NET audit cycles on isolated synthetic data, not production UAT."""
from copy import deepcopy
from datetime import datetime,timedelta,timezone
from decimal import Decimal
from hashlib import sha256
from io import BytesIO
import json
from zipfile import ZipFile
import pytest
from openpyxl import load_workbook
from sqlalchemy import select,func
from sqlalchemy.orm import Session
from tests.api.conftest import api_factory,synthetic_password_hash
from tests.features.conftest import features_api
from tests.workspace_v2.test_access_and_registry import workspace,grants
from tests.workspace_v2.test_linked_references import linked
from tests.workspace_v2.invoice_fixtures import attach_invoices
from tests.workspace_v2.net_cycle_fixtures import (case_row,import_cases,configure_signers,
    approve_batch_policies,execute_batch,result_files)
from mcias.models import SourceFile,AuditLog
from mcias.models_claims import ClaimEntry
from mcias.models_ingestion import IngestionDocument
from mcias.models_audit import AuditSnapshot
from mcias.models_audit_policy import AuditPolicyGrant
from mcias.services.registry_service import store
from mcias.palermo2 import verify
from mcias.seed.synthetic_invoices import invoice_pdf


@pytest.fixture
def cycle(linked,monkeypatch):
    configure_signers(linked,monkeypatch);return linked


def assert_source_unchanged(api,batch):
    with Session(api.engine) as db:
        source=db.get(SourceFile,db.get(IngestionDocument,batch.document_id).source_id)
        assert source.sha256==sha256(batch.raw).hexdigest()
        assert store().get(source.storage_key,source.sha256)==batch.raw


def test_one_complete_cycle_with_real_invoice_policy_snapshot_approval_and_export(cycle):
    api=cycle;batch=import_cases(api);approve_batch_policies(api,batch.batch_id)
    run=execute_batch(api,batch.batch_id)
    assert run['counts']=={'accepted':1,'rejected':0,'held':0,'review':0},run
    row=run['rows'][0]
    with Session(api.engine) as db:
        snapshot=db.get(AuditSnapshot,row['snapshot_id']);claim=db.get(ClaimEntry,row['line_id'])
        assert snapshot.verification['status']=='verified'
        assert snapshot.calculation['amount_basis']=='NET'
        assert Decimal(snapshot.calculation['claimed'])==Decimal('900')
        assert Decimal(snapshot.calculation['payable'])==Decimal('900')
        assert Decimal(snapshot.calculation['gross_evidence']['claimed'])==Decimal('1000')
        assert snapshot.calculation['copay_deducted_again'] is False
        assert not any(x['rule']=='copay' for x in snapshot.calculation['deductions'])
        assert verify(claim,snapshot.findings['agents'],snapshot.calculation)['status']=='verified'
        source=db.get(SourceFile,db.get(IngestionDocument,batch.invoice_document_id).source_id)
        assert store().get(source.storage_key,source.sha256)==batch.invoice_raw
    case=api.client.post('/api/v1/governance/cases/from-audit/'+row['snapshot_id'],headers=api.headers['root'])
    assert case.status_code==201,case.text
    case_id=case.json()['id'];path='/api/v1/governance/cases/'+case_id+'/events'
    proposal=api.client.post(path,json={'kind':'propose','payload':{'reason_code':'SYN_REVIEWED'}},headers=api.headers['root'])
    assert proposal.status_code==201,proposal.text
    assert api.client.post(path,json={'kind':'approve','payload':{}},headers=api.headers['root']).status_code==403
    decision=api.client.post(path,json={'kind':'approve','payload':{'reason_code':'SYN_APPROVED'}},headers=api.headers['none'])
    assert decision.status_code==201,decision.text
    history=api.client.get('/api/v1/governance/cases/'+case_id,headers=api.headers['root']).json()
    assert [event['kind'] for event in history['events']]==['propose','approve']
    final_report=api.client.post('/api/v1/reports/batches/'+batch.batch_id+'/snapshots',headers=api.headers['root'])
    assert final_report.status_code==201,final_report.text
    money=final_report.json()['summary']['by_currency']['LYD']
    assert money['approved_count']==1 and Decimal(money['payable'])==Decimal('900')
    assert Decimal(money['deductions'])==0
    final_excel=api.client.get('/api/v1/reports/snapshots/'+final_report.json()['id']+'/excel?lang=ar',headers=api.headers['root'])
    assert final_excel.status_code==200 and final_excel.content.startswith(b'PK'),final_excel.text[:200]
    result,raw=result_files(api,batch.batch_id)
    assert result['summary']['counts']==run['counts'] and result['financially_final'] is False
    with ZipFile(BytesIO(raw)) as archive:
        assert set(archive.namelist())=={'accepted.xlsx','rejected.xlsx','held.xlsx','review.xlsx'}
        assert load_workbook(BytesIO(archive.read('accepted.xlsx')),read_only=True).active.max_row==2
        assert all(load_workbook(BytesIO(archive.read(key+'.xlsx')),read_only=True).active.max_row==1 for key in ('rejected','held','review'))
    assert_source_unchanged(api,batch)


def test_multi_case_xlsx_has_one_exclusive_file_per_service_line(cycle):
    api=cycle;cases=[case_row(1),case_row(2,gross='1200',unit='1200',copay='120',net='1080'),
        case_row(3,name='SYN wrong identity'),case_row(4)]
    batch=import_cases(api,cases,missing_references={cases[3]['invoice_reference']})
    approve_batch_policies(api,batch.batch_id);run=execute_batch(api,batch.batch_id)
    assert run['counts']=={'accepted':1,'rejected':1,'held':1,'review':1},run
    by_id={r['line_id']:r for r in run['rows']}
    assert by_id[batch.identities[cases[3]['invoice_reference']]]['reason']=='missing_invoice'
    result,raw=result_files(api,batch.batch_id);assert result['summary']['counts']==run['counts']
    ids=[]
    with ZipFile(BytesIO(raw)) as archive:
        for name in archive.namelist():
            rows=list(load_workbook(BytesIO(archive.read(name)),read_only=True).active.values)
            assert len(rows)==2
            ids.append(rows[1][0])
            if name=='held.xlsx':assert 'نقص المرفقات' in str(rows[1][9])
    assert len(ids)==len(set(ids))==len(cases)
    assert set(ids)==set(batch.identities.values())
    rejected=by_id[batch.identities[cases[1]['invoice_reference']]]
    with Session(api.engine) as db:
        calc=db.get(AuditSnapshot,rejected['snapshot_id']).calculation
        assert Decimal(calc['claimed'])==Decimal('1080')
        assert Decimal(calc['payable'])==Decimal('900')
        assert Decimal(calc['total_deduction'])==Decimal('180')
    assert_source_unchanged(api,batch)


def test_missing_invoice_is_held_then_recovers_after_actual_invoice_arrives(cycle):
    api=cycle;batch=import_cases(api,attach=False);approve_batch_policies(api,batch.batch_id)
    held=execute_batch(api,batch.batch_id)
    assert held['counts']['held']==1 and held['rows'][0]['reason']=='missing_invoice'
    with Session(api.engine) as db:assert db.scalar(select(func.count()).select_from(AuditSnapshot))==0
    attach_invoices(api,batch.batch_id)
    first=execute_batch(api,batch.batch_id);again=execute_batch(api,batch.batch_id)
    assert first['counts']['accepted']==1 and first['rows'][0]['snapshot_id']==again['rows'][0]['snapshot_id']
    assert first['fingerprint']==again['fingerprint']
    with Session(api.engine) as db:
        assert db.scalar(select(func.count()).select_from(AuditSnapshot))==1
        assert db.scalar(select(func.count()).select_from(AuditLog).where(AuditLog.action=='net_audit.completed'))==2


def test_missing_invoice_overrides_other_review_findings(cycle):
    api=cycle;batch=import_cases(api,[case_row(name='SYN wrong identity')],attach=False)
    approve_batch_policies(api,batch.batch_id);run=execute_batch(api,batch.batch_id)
    assert run['counts']=={'accepted':0,'rejected':0,'held':1,'review':0}
    assert run['rows'][0]['reason']=='missing_invoice'
    result,_=result_files(api,batch.batch_id);assert result['summary']['counts']==run['counts']


def test_invoice_is_required_for_each_line_not_just_anywhere_in_batch(cycle):
    cases=[case_row(1),case_row(2)]
    batch=import_cases(cycle,cases,missing_references={cases[1]['invoice_reference']})
    approve_batch_policies(cycle,batch.batch_id);run=execute_batch(cycle,batch.batch_id)
    assert run['counts']=={'accepted':1,'rejected':0,'held':1,'review':0}


def test_original_xlsx_cannot_be_mislabeled_as_invoice(cycle):
    api=cycle;batch=import_cases(api,attach=False);identity=next(iter(batch.identities.values()))
    response=api.client.post(f'/api/v1/claims/entries/{identity}/attachments/{batch.document_id}?kind=invoice',headers=api.headers['root'])
    assert response.status_code==201,response.text
    approve_batch_policies(api,batch.batch_id);run=execute_batch(api,batch.batch_id)
    assert run['rows'][0]['reason']=='missing_invoice'


def test_supporting_pdf_is_not_automatically_an_invoice(cycle):
    api=cycle;batch=import_cases(api,attach=False);identity=next(iter(batch.identities.values()))
    document,_=attach_invoices(api,batch.batch_id,skip_ids={identity})
    response=api.client.post(f'/api/v1/claims/entries/{identity}/attachments/{document}?kind=supporting',headers=api.headers['root'])
    assert response.status_code==201,response.text
    approve_batch_policies(api,batch.batch_id);run=execute_batch(api,batch.batch_id)
    assert run['rows'][0]['reason']=='missing_invoice'


def test_signed_wrong_copay_cannot_override_category_reference(cycle):
    api=cycle;batch=import_cases(api);approve_batch_policies(api,batch.batch_id,overrides={'copay_percent':'50'})
    run=execute_batch(api,batch.batch_id)
    assert run['counts']['held']==1 and run['rows'][0]['reason']=='policy_or_evidence_blocked'
    assert 'category' in run['rows'][0]['detail']
    with Session(api.engine) as db:assert db.scalar(select(func.count()).select_from(AuditSnapshot))==0


def test_unsigned_policy_and_technical_financial_approval_are_rejected(cycle):
    api=cycle;batch=import_cases(api)
    requirements=api.client.get(f'/api/v1/audits/batches/{batch.batch_id}/policy-requirements',headers=api.headers['root']).json()
    calc=requirements['items'][0]['reference_calculation']
    definition={'policy_id':'SYN-UNSIGNED','currency':calc['currency'],'rounding':calc['rounding'],
        'precision':calc['precision'],'order':['exclusion','tariff','cap','copay'],'copay_percent':'10',
        'cap':None,'exclusion':'0','tariff':'1000','approved':True,'amount_basis':'NET',
        'reference_fingerprint':requirements['items'][0]['reference_fingerprint']}
    assert api.client.post('/api/v1/audit-policies/definitions',json=definition,headers=api.headers['root']).status_code==403
    assert api.client.post('/api/v1/audit-policies/definitions',json=definition,headers=api.headers['none']).status_code==409
    assert api.client.post(f'/api/v1/audits/batches/{batch.batch_id}/run').status_code==401
    assert api.client.post(f'/api/v1/audits/batches/{batch.batch_id}/run',headers=api.headers['other']).status_code==403
    run=execute_batch(api,batch.batch_id);assert run['rows'][0]['reason']=='missing_or_ambiguous_signed_policy'


def test_palermo_recomputes_net_after_attacker_reseals_wrong_payment(cycle):
    api=cycle;batch=import_cases(api);approve_batch_policies(api,batch.batch_id);run=execute_batch(api,batch.batch_id)
    with Session(api.engine) as db:
        snapshot=db.get(AuditSnapshot,run['rows'][0]['snapshot_id']);claim=db.get(ClaimEntry,run['rows'][0]['line_id'])
        calc=deepcopy(snapshot.calculation);calc.update(payable='800.00',total_deduction='100.00')
        calc['deductions']=[{'rule':'net_reference_adjustment','deduction':'100.00','remaining':'800.00'}]
        calc['fingerprint']=sha256(json.dumps({k:v for k,v in calc.items() if k!='fingerprint'},sort_keys=True,separators=(',',':')).encode()).hexdigest()
        verification=verify(claim,snapshot.findings['agents'],calc)
        assert verification['status']=='held' and 'NET_payable_mismatch' in verification['issues']


def test_revoked_policy_does_not_reuse_old_snapshot_as_new_accepted_run(cycle):
    api=cycle;batch=import_cases(api);policies=approve_batch_policies(api,batch.batch_id)
    assert execute_batch(api,batch.batch_id)['counts']['accepted']==1
    policy=next(iter(policies.values()))
    revoked=api.client.post('/api/v1/audit-policies/grants/'+policy['grant_id']+'/revoke',headers=api.headers['none'])
    assert revoked.status_code==200,revoked.text
    assert execute_batch(api,batch.batch_id)['counts']['held']==1
    result,_=result_files(api,batch.batch_id);assert result['summary']['counts']['held']==1
    with Session(api.engine) as db:assert db.scalar(select(func.count()).select_from(AuditSnapshot))==1


def reseal(calculation):
    """Model an attacker who can rewrite producer evidence and its plain hash."""
    calculation['fingerprint']=sha256(json.dumps(
        {key:value for key,value in calculation.items() if key!='fingerprint'},
        sort_keys=True,separators=(',',':')).encode()).hexdigest()
    return calculation


def bound_verification(claim,snapshot,calculation,reference_check=None):
    return verify(claim,snapshot.findings['agents'],calculation,
        policy_fingerprint=snapshot.policy_fingerprint,policy_id=snapshot.policy_id,
        reference_check=snapshot.findings['reference_check'] if reference_check is None else reference_check)


def test_palermo_signed_anchor_rejects_resealed_self_consistent_copay(cycle):
    api=cycle;batch=import_cases(api);approve_batch_policies(api,batch.batch_id)
    run=execute_batch(api,batch.batch_id)
    with Session(api.engine) as db:
        snapshot=db.get(AuditSnapshot,run['rows'][0]['snapshot_id'])
        claim=db.get(ClaimEntry,run['rows'][0]['line_id'])
        assert bound_verification(claim,snapshot,snapshot.calculation)['status']=='verified'
        calc=deepcopy(snapshot.calculation)
        calc['policy_definition']['copay_percent']='50'
        calc['policy_fingerprint']=sha256(json.dumps(calc['policy_definition'],
            sort_keys=True,separators=(',',':')).encode()).hexdigest()
        gross=calc['gross_evidence']
        gross['deductions'][-1]={'rule':'copay','deduction':'500.00','remaining':'500.00'}
        gross.update(payable='500.00',total_deduction='500.00');reseal(gross)
        calc.update(payable='500.00',total_deduction='400.00')
        calc['deductions']=[{'rule':'net_reference_adjustment','deduction':'400.00','remaining':'500.00'}]
        reseal(calc)
        # Internal arithmetic and producer hashes agree, but the actual signed
        # institutional policy remains the original category B copay of 10%.
        assert verify(claim,snapshot.findings['agents'],calc)['status']=='verified'
        checked=bound_verification(claim,snapshot,calc)
        assert checked['status']=='held' and 'NET_signed_policy_mismatch' in checked['issues']


def test_palermo_confirmed_mapping_cannot_be_replaced_by_equal_value_decoy(cycle):
    api=cycle;batch=import_cases(api);approve_batch_policies(api,batch.batch_id)
    run=execute_batch(api,batch.batch_id)
    with Session(api.engine) as db:
        snapshot=db.get(AuditSnapshot,run['rows'][0]['snapshot_id'])
        claim=db.get(ClaimEntry,run['rows'][0]['line_id'])
        calc=deepcopy(snapshot.calculation)
        assert calc['source_columns']['gross_amount']=='المطلوب'
        assert claim.source_values['المطلوب']==claim.source_values['السعر']
        calc['source_columns']['gross_amount']='السعر';reseal(calc)
        assert verify(claim,snapshot.findings['agents'],calc)['status']=='verified'
        checked=bound_verification(claim,snapshot,calc)
        assert checked['status']=='held' and 'NET_confirmed_mapping_mismatch' in checked['issues']


def test_palermo_resealed_wrong_ids_and_malformed_reference_bindings_fail_closed(cycle):
    api=cycle;batch=import_cases(api);approve_batch_policies(api,batch.batch_id)
    run=execute_batch(api,batch.batch_id)
    with Session(api.engine) as db:
        snapshot=db.get(AuditSnapshot,run['rows'][0]['snapshot_id'])
        claim=db.get(ClaimEntry,run['rows'][0]['line_id'])
        for field,issue in (('policy_id','NET_policy_id_mismatch'),
            ('contract_id','NET_contract_binding_mismatch'),
            ('enrollment_id','NET_enrollment_binding_mismatch'),
            ('source_id','source_mismatch')):
            calc=deepcopy(snapshot.calculation);calc[field]='SYN-UNRELATED-ID';reseal(calc)
            checked=bound_verification(claim,snapshot,calc)
            assert checked['status']=='held' and issue in checked['issues'],(field,checked)
        for malformed in ([],{}, {'references':[]},{'references':{'contract':None}}):
            checked=bound_verification(claim,snapshot,snapshot.calculation,malformed)
            assert checked['status']=='held'
            assert {'NET_reference_check_missing','NET_invalid_reference_binding'}&set(checked['issues'])


def test_palermo_cannot_downgrade_trusted_net_proof_to_gross(cycle):
    api=cycle;batch=import_cases(api);approve_batch_policies(api,batch.batch_id)
    run=execute_batch(api,batch.batch_id)
    with Session(api.engine) as db:
        snapshot=db.get(AuditSnapshot,run['rows'][0]['snapshot_id'])
        claim=db.get(ClaimEntry,run['rows'][0]['line_id'])
        calc=deepcopy(snapshot.calculation);calc['amount_basis']='GROSS';reseal(calc)
        checked=bound_verification(claim,snapshot,calc)
        assert checked['status']=='held' and 'NET_amount_basis_mismatch' in checked['issues']


def test_changed_invoice_evidence_requires_reaudit_and_keeps_old_frozen_output(cycle):
    api=cycle;batch=import_cases(api);approve_batch_policies(api,batch.batch_id)
    run=execute_batch(api,batch.batch_id);assert run['counts']['accepted']==1
    original,original_zip=result_files(api,batch.batch_id)
    line_id=run['rows'][0]['line_id']
    corrected=invoice_pdf([{'invoice_reference':'SYN-CORRECTED-INVOICE-1',
        'beneficiary':'SYN-INS-EMPLOYEE','provider':'SYN-LINKED-HOSPITAL',
        'service':'SYN-VISIT','service_date':batch.cases[0]['date'],
        'gross':'1000','copay':'100','net':'900','currency':'LYD'}])
    assert sha256(corrected).digest()!=sha256(batch.invoice_raw).digest()
    upload=api.client.post('/api/v1/ingestion/uploads?filename=SYN-corrected-invoice.pdf',
        content=corrected,headers={**api.headers['root'],'Content-Type':'application/octet-stream'})
    assert upload.status_code==200 and upload.json()['status']=='extracted',upload.text
    linked=api.client.post(f"/api/v1/claims/entries/{line_id}/attachments/{upload.json()['document_id']}?kind=invoice",
        headers=api.headers['root'])
    assert linked.status_code==201,linked.text
    updated,_=result_files(api,batch.batch_id)
    assert updated['id']!=original['id']
    assert updated['summary']['counts']=={'accepted':0,'rejected':0,'held':0,'review':1}
    lines=api.client.get('/api/v1/result-files/'+updated['id']+'/lines',headers=api.headers['root'])
    assert lines.status_code==200,lines.text
    row=lines.json()['items'][0]
    assert row['reason']=='invoice_evidence_changed_reaudit_required' and row['proposed_net'] is None
    old=api.client.get('/api/v1/result-files/'+original['id']+'/lines',headers=api.headers['root'])
    assert old.status_code==200 and old.json()['items'][0]['bucket']=='accepted'
    download=api.client.get('/api/v1/result-files/'+original['id']+'/download',headers=api.headers['root'])
    assert download.status_code==200
    with ZipFile(BytesIO(download.content)) as archive,ZipFile(BytesIO(original_zip)) as original_archive:
        for name in archive.namelist():
            assert list(load_workbook(BytesIO(archive.read(name)),read_only=True).active.values)==list(
                load_workbook(BytesIO(original_archive.read(name)),read_only=True).active.values)
    assert_source_unchanged(api,batch)


def test_financial_export_requires_claim_source_read_permission(cycle):
    api=cycle;batch=import_cases(api);approve_batch_policies(api,batch.batch_id)
    execute_batch(api,batch.batch_id);result,_=result_files(api,batch.batch_id)
    assert grants(api,['reports.read','reports.export'],user='admin').status_code==200
    path='/api/v1/result-files/'+result['id']
    assert api.client.get('/api/v1/result-files',headers=api.headers['admin']).status_code==200
    assert api.client.get(path+'/lines',headers=api.headers['admin']).status_code==403
    assert api.client.get(path+'/download',headers=api.headers['admin']).status_code==403
    assert grants(api,['reports.read','reports.export','claims.read'],user='admin').status_code==200
    assert api.client.get(path+'/lines',headers=api.headers['admin']).status_code==200
    assert api.client.get(path+'/download',headers=api.headers['admin']).status_code==200


def test_policy_revocation_is_one_way_even_when_orm_attribute_is_expired(cycle):
    api=cycle;batch=import_cases(api);policies=approve_batch_policies(api,batch.batch_id)
    policy=next(iter(policies.values()));grant_id=policy['grant_id']
    with Session(api.engine) as db:
        grant=db.get(AuditPolicyGrant,grant_id)
        original_policy_id=grant.policy_id
        assert grant.revoked_at is None
        # A first revocation must not be a vehicle for rewriting its approval.
        grant.policy_id='SYN-UNAUTHORIZED-POLICY';grant.revoked_at=datetime.now(timezone.utc)
        with pytest.raises(ValueError,match='approval identity is immutable'):db.flush()
        db.rollback()
        assert db.get(AuditPolicyGrant,grant_id).policy_id==original_policy_id
        assert db.get(AuditPolicyGrant,grant_id).revoked_at is None
    path='/api/v1/audit-policies/grants/'+grant_id+'/revoke'
    first=api.client.post(path,headers=api.headers['none'])
    assert first.status_code==200 and first.json()['revoked'] is True,first.text
    second=api.client.post(path,headers=api.headers['none'])
    assert second.status_code==403 and 'already revoked' in second.text,second.text
    with Session(api.engine) as db:
        grant=db.get(AuditPolicyGrant,grant_id);original_revoked_at=grant.revoked_at
        assert original_revoked_at is not None
        for expired in (False,True):
            for attempted in (None,datetime.now(timezone.utc)+timedelta(days=1)):
                if expired:db.expire(grant,['revoked_at'])
                grant.revoked_at=attempted
                with pytest.raises(ValueError,match='only first revocation is allowed'):db.flush()
                db.rollback()
                assert grant.revoked_at==original_revoked_at
        db.expire(grant,['revoked_at'])
        grant.policy_id='SYN-UNAUTHORIZED-POLICY'
        with pytest.raises(ValueError):db.flush()
        db.rollback()
        assert grant.policy_id==original_policy_id and grant.revoked_at==original_revoked_at
