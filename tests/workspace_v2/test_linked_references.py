"""Real linked inputs -> actual XLSX import -> service-date pre-audit checks."""
from datetime import date, timedelta
from hashlib import sha256
from io import BytesIO
import pytest
from openpyxl import load_workbook
from sqlalchemy import select, func
from sqlalchemy.orm import Session
from tests.api.conftest import api_factory, synthetic_password_hash
from tests.features.conftest import features_api
from tests.workspace_v2.test_access_and_registry import workspace, grants
from tests.workspace_v2.test_mapping_and_results import mapping, upload_batch
from mcias.models import User, AuditLog, SourceFile
from mcias.models_ingestion import IngestionDocument
from mcias.models_claims import ClaimEntry
from mcias.seed.linked_references import seed_linked_references, sample_workbook
from mcias.services import registry_service as registry


@pytest.fixture
def linked(workspace,monkeypatch):
    api=workspace
    monkeypatch.setenv('MCIAS_ENVIRONMENT','development');monkeypatch.setenv('MCIAS_LAPTOP_TEST_MODE','synthetic_localhost')
    with Session(api.engine) as db:
        result=seed_linked_references(db,db.get(User,api.ids['root']));db.commit()
    api.linked=result['records'];return api


def linked_mapping():
    payload=mapping();payload['columns'].update({'beneficiary_name':'اسم المستفيد','category_code':'الدرجة',
        'gross_amount':'المطلوب','unit_price':'السعر','copay_amount':'إجمالي التحمل','copay_percent':'نسبة التحمل'})
    return payload


def import_sample(api,updates=None,at=None):
    raw=sample_workbook(at)
    if updates:
        book=load_workbook(BytesIO(raw));sheet=book.active
        for (row,header),value in updates.items():
            columns={c.value:c.column for c in sheet[1]};sheet.cell(row,columns[header],value)
        output=BytesIO();book.save(output);raw=output.getvalue()
    batch,doc=upload_batch(api,raw,'SYN-linked-claims.xlsx')
    base='/api/v1/claims/batches/'+batch
    response=api.client.put(base+'/mapping',json=linked_mapping(),headers=api.headers['root'])
    assert response.status_code==200,response.text
    response=api.client.post(base+'/resume',headers=api.headers['root'])
    assert response.status_code==200,response.text
    assert response.json()['reference_check']['summary']['total_lines']==4
    from tests.workspace_v2.invoice_fixtures import attach_invoices
    attach_invoices(api,batch)
    assert api.client.post(base+'/reference-checks',headers=api.headers['root']).status_code==201
    result=api.client.get(base+'/reference-checks',headers=api.headers['root'])
    assert result.status_code==200,result.text
    return result.json(),batch,doc,raw


def first(api,result):
    with Session(api.engine) as db:
        identities=set(db.scalars(select(ClaimEntry.id).where(ClaimEntry.organization_id==api.organization_id,
            ClaimEntry.batch_id==result['batch_id'],ClaimEntry.claim_reference=='SYN-INVOICE-1',ClaimEntry.service_code=='SYN زيارة طبيب')))
    return next(r for r in result['rows'] if r['line_id'] in identities)


def test_linked_seed_is_idempotent_and_keeps_financial_approvals_closed(linked):
    api=linked
    with Session(api.engine) as db:
        again=seed_linked_references(db,db.get(User,api.ids['root']));db.commit()
        assert again['records_created']==0 and len(again['records'])==50
        assert all(row['financially_active'] is False for row in again['records'].values())
    assert api.client.get('/api/v1/me/features',headers=api.headers['admin']).json()['permissions']==[]


def test_category_service_network_and_provider_price_are_resolved_automatically(linked):
    result,_,_,_=import_sample(linked)
    assert result['summary']=={'total_lines':4,'MATCHED':4,'REVIEW_REQUIRED':0,'BLOCKED':0}
    calculations=[r['calculation'] for r in result['rows']]
    values={(c['category_code'],c['service_type'],c['network']):(c['beneficiary_copay_percent'],c['expected_net']) for c in calculations}
    assert values=={('B','visit','in_network'):('10','900.00'),('B','lab','in_network'):('30','140.00'),
        ('C','visit','in_network'):('50','500.00'),('B','visit','out_of_network'):('30','840.00')}
    assert all(c['copay_deducted_again'] is False for c in calculations)
    assert not result['audit_completed'] and not result['financially_final']
    assert all({'contract','regulation','beneficiary_enrollment','category','provider','service','provider_tariff','copay_rule'}<=set(r['references']) for r in result['rows'])


@pytest.mark.parametrize('updates,code',[
    ({(2,'اسم المستفيد'):'SYN شخص آخر'},'beneficiary_name'),
    ({(2,'الدرجة'):'C'},'source_category'),
    ({(2,'نسبة التحمل'):'50'},'copay_percent'),
    ({(2,'إجمالي التحمل'):'200'},'copay_amount'),
    ({(2,'الصافي'):'800'},'net_amount'),
    ({(2,'المطلوب'):'1200',(2,'السعر'):'1200',(2,'إجمالي التحمل'):'120',(2,'الصافي'):'1080'},'price'),
    ({(2,'المطلوب'):'1100'},'gross_total'),
])
def test_source_discrepancies_require_review_without_changing_original(linked,updates,code):
    result,_,doc,raw=import_sample(linked,updates)
    row=first(linked,result)
    assert row['status']=='REVIEW_REQUIRED'
    assert any(c['code']==code and c['status']=='difference' for c in row['checks'])
    with Session(linked.engine) as db:
        source=db.get(SourceFile,db.get(IngestionDocument,doc).source_id)
        assert source.sha256==sha256(raw).hexdigest() and registry.store().get(source.storage_key,source.sha256)==raw


@pytest.mark.parametrize('updates',[
    {(2,'الرقم التأميني'):'SYN-UNKNOWN'}, {(2,'مقدم الخدمة'):'SYN-UNKNOWN'},
    {(2,'الخدمة'):'SYN-UNKNOWN'}, {(2,'المطلوب'):'not-a-number'},
    {(2,'المطلوب'):'1000%'}, {(2,'المطلوب'):'1,00'},
])
def test_missing_or_invalid_sources_are_held_not_financially_rejected(linked,updates):
    result,_,_,_=import_sample(linked,updates)
    row=next(r for r in result['rows'] if r['status']=='BLOCKED')
    assert not row['financially_final'] and any(c['status']=='hold' for c in row['checks'])


def test_service_date_outside_subscription_does_not_use_current_category(linked):
    result,_,_,_=import_sample(linked,at=date(date.today().year-1,12,31))
    assert result['summary']['BLOCKED']==4 and result['summary']['MATCHED']==0


def test_missing_tariff_is_not_invented_from_invoice(linked):
    api=linked;price=api.linked['PRICE-HOSPITAL-visit']
    path='/api/v1/registry/providers.prices/records/'+price['id']+'/close'
    assert api.client.post(path,json={'expected_version':price['version'],'reason':'SYN missing tariff test'},headers=api.headers['root']).status_code==200
    result,_,_,_=import_sample(api)
    row=first(api,result)
    assert row['status']=='BLOCKED' and row['calculation'] is None


def test_bad_cross_tenant_and_inconsistent_reference_links_are_rejected(linked):
    api=linked;rate=api.linked['COPAY-HOSPITAL-B-visit']
    from mcias.models import Role, UserRole, MFAFactor
    from mcias.schemas.registry import RecordCreate
    from mcias.services import feature_service as features
    with Session(api.engine) as db:
        other=db.get(User,api.ids['other'])
        role=db.scalar(select(Role).where(Role.code=='super_admin'))
        db.add(UserRole(user_id=other.id,role_id=role.id))
        db.add(MFAFactor(user_id=other.id,enabled=True,encrypted_secret=b'SYN-unused-no-login-secret'))
        db.flush()
        features.grant_feature(db,other.organization_id,'beneficiaries.categories',other.id)
        foreign=registry.create_record(db,other,'beneficiaries.categories',RecordCreate(title='SYN foreign category',
            effective_from=date.today(),values={'record_code':'SYN-FOREIGN-CATEGORY','category_code':'B'}))
        foreign_id=foreign['id'];db.commit()
    data={key:rate[key] for key in ('title','effective_from','values','links','terms')}
    data['values']={**data['values'],'record_code':'SYN-BAD-LINK'}
    data['links']={**data['links'],'category':foreign_id}
    response=api.client.post('/api/v1/registry/beneficiaries.cost_sharing/records',json=data,headers=api.headers['root'])
    assert response.status_code==409
    price=api.linked['PRICE-HOSPITAL-visit'];data={key:price[key] for key in ('title','effective_from','values','links','terms')}
    data['values']={**data['values'],'record_code':'SYN-WRONG-PROVIDER'}
    data['links']={**data['links'],'provider':api.linked['CLINIC']['id']}
    assert api.client.post('/api/v1/registry/providers.prices/records',json=data,headers=api.headers['root']).status_code==409


def test_duplicate_linked_copay_cannot_silently_compete(linked):
    api=linked;rate=api.linked['COPAY-HOSPITAL-B-visit']
    data={key:rate[key] for key in ('title','effective_from','values','links','terms')}
    data['values']={**data['values'],'record_code':'SYN-DUPLICATE-COPAY','percent':'50'}
    assert api.client.post('/api/v1/registry/beneficiaries.cost_sharing/records',json=data,headers=api.headers['root']).status_code==409


def test_future_rule_does_not_change_current_or_historical_copay(linked):
    api=linked;rate=api.linked['COPAY-HOSPITAL-B-visit'];future=date.today()+timedelta(days=30)
    data={key:rate[key] for key in ('title','values','links','terms')}
    data.update({'effective_from':future.isoformat(),'expected_version':rate['version'],'reason':'SYN future copay'})
    data['values']={**data['values'],'percent':'20'}
    response=api.client.post('/api/v1/registry/beneficiaries.cost_sharing/records/'+rate['id']+'/replace',json=data,headers=api.headers['root'])
    assert response.status_code==201,response.text
    current,_,_,_=import_sample(api)
    assert first(api,current)['calculation']['beneficiary_copay_percent']=='10'
    later,_,_,_=import_sample(api,at=future)
    assert first(api,later)['calculation']['beneficiary_copay_percent']=='20'
    assert first(api,later)['status']=='REVIEW_REQUIRED'


def test_check_is_idempotent_tenant_scoped_and_permission_guarded(linked):
    result,batch,_,_=import_sample(linked)
    api=linked;base='/api/v1/claims/batches/'+batch+'/reference-checks'
    again=api.client.post(base,headers=api.headers['root'])
    assert again.status_code==201 and again.json()['id']==result['id']
    assert api.client.get(base,headers=api.headers['other']).status_code==403
    assert grants(api,['claims.read']).status_code==200
    assert api.client.get(base,headers=api.headers['admin']).status_code==403
    assert grants(api,['claims.read','claims.review']).status_code==200
    assert api.client.get(base,headers=api.headers['admin']).status_code==200
    with Session(api.engine) as db:
        assert db.scalar(select(func.count()).select_from(AuditLog).where(AuditLog.action=='claim_references.checked',AuditLog.resource_id==batch))==2


def test_legacy_gross_evaluator_cannot_apply_copay_again_to_imported_net(linked):
    from mcias.audit_service import _require_supported_amount_basis
    from mcias.audit_engine import PolicyBlocked
    _,batch,_,_=import_sample(linked)
    with Session(linked.engine) as db:
        claim=db.scalar(select(ClaimEntry).where(ClaimEntry.batch_id==batch))
        with pytest.raises(PolicyBlocked,match='must not deduct copay again'):
            _require_supported_amount_basis(db,claim)


def test_expired_replacement_does_not_reactivate_an_older_reference(linked):
    from mcias.schemas.registry import RecordReplace
    api=linked;rate=api.linked['COPAY-HOSPITAL-B-visit']
    tomorrow=date.today()+timedelta(days=1);end=tomorrow+timedelta(days=2)
    data={key:rate[key] for key in ('title','values','links','terms')}
    data.update(effective_from=tomorrow,effective_to=end,expected_version=rate['version'],reason='SYN limited replacement')
    with Session(api.engine) as db:
        root=db.get(User,api.ids['root'])
        registry.replace_record(db,root,'beneficiaries.cost_sharing',rate['id'],RecordReplace(**data));db.commit()
        ids={r['id'] for r in registry.effective_records(db,root,'beneficiaries.cost_sharing',end)}
        assert rate['id'] not in ids


def test_amount_grouping_cannot_silently_change_net_and_percent_is_not_a_price():
    from mcias.claim_normalization import source_decimal, ClaimValidationError
    assert source_decimal('1,000.00')==1000
    assert source_decimal('10%',allow_percent=True)==10
    for value in ('9,00','1000%','NaN','-1'):
        with pytest.raises(ClaimValidationError):source_decimal(value)
