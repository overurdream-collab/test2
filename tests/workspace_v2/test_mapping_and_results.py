"""Real upload/import/export; synthetic pre-existing G5 evidence, not agent acceptance."""
from datetime import date
from decimal import Decimal
from hashlib import sha256
from io import BytesIO
from zipfile import ZipFile
import json
import pytest
from openpyxl import load_workbook, Workbook
from sqlalchemy import select
from sqlalchemy.orm import Session
from tests.api.conftest import api_factory, synthetic_password_hash
from tests.features.conftest import features_api
from tests.workspace_v2.test_access_and_registry import workspace
from mcias.models import SourceFile
from mcias.models_claims import ClaimEntry, ClaimBatch
from mcias.models_ingestion import IngestionDocument
from mcias.models_audit import AuditSnapshot
from mcias.audit_agents import AGENTS
from mcias.audit_engine import AuditPolicy, evaluate_net
from mcias.reporting_g7 import canonical, ReportBlocked
from mcias.services import result_service as results
from mcias.services.registry_service import store


def upload_batch(api,raw,filename='SYN-claims.xlsx'):
    upload=api.client.post('/api/v1/ingestion/uploads?filename='+filename,content=raw,
        headers={**api.headers['root'],'Content-Type':'application/octet-stream'})
    assert upload.status_code==200,upload.text
    doc=upload.json()['document_id']
    stage=api.client.post('/api/v1/claims/batches/from-document/'+doc,headers=api.headers['root'])
    assert stage.status_code==201,stage.text
    return stage.json()['id'],doc


def arabic_workbook():
    book=Workbook();sheet=book.active;sheet.title='تفاصيل المطالبات'
    sheet.append(['رقم الفاتورة','الرقم التأميني','مقدم الخدمة','الخدمة','تاريخ الخدمة','المطلوب','إجمالي التحمل','الصافي'])
    for service in ('SYN-visit','SYN-test','SYN-drug','SYN-xray'):
        sheet.append(['SYN-SAME-INVOICE','SYN-BENEFICIARY','SYN-HOSPITAL',service,'08/06/2026','100','20','80'])
    stream=BytesIO();book.save(stream);return stream.getvalue()


def mapping():
    return {'columns':{'claim_reference':'رقم الفاتورة','beneficiary_token':'الرقم التأميني',
        'provider_code':'مقدم الخدمة','service_code':'الخدمة','service_date':'تاريخ الخدمة','claimed_amount':'الصافي'},
        'currency':'LYD','date_format':'DMY'}


def imported(api):
    raw=arabic_workbook();batch,doc=upload_batch(api,raw)
    path='/api/v1/claims/batches/'+batch
    mapped=api.client.put(path+'/mapping',json=mapping(),headers=api.headers['root'])
    assert mapped.status_code==200,mapped.text
    response=api.client.post(path+'/resume',headers=api.headers['root'])
    assert response.status_code==200 and response.json()['imported_count']==4,response.text
    return batch,doc,raw


def snapshots(api,batch):
    from tests.workspace_v2.invoice_fixtures import attach_invoices
    attach_invoices(api,batch)
    with Session(api.engine) as db:
        claims=list(db.scalars(select(ClaimEntry).where(ClaimEntry.batch_id==batch).order_by(ClaimEntry.service_code)))
        for i,claim in enumerate(claims):
            findings=[{'agent':name,'status':'pass','reason':'SYN checked evidence','reference':'SYN-reference'} for name in AGENTS]
            reference_check=None;policy_fp='SYN-fingerprint'
            if i==2:calc={'status':'held','reason':'missing_reference'};verification={'status':'held'}
            else:
                from mcias.services.invoice_service import invoices
                from mcias.services.net_audit_service import reference_fingerprint
                from mcias.audit_service import fingerprint_policy
                reference_check={'calculation':{'quantity':'1','gross':'100'},
                    'references':{'contract':{'id':'SYN-contract'},'beneficiary_enrollment':{'id':'SYN-enrollment'}},
                    'invoices':invoices(db,claim),'source_columns':{**mapping()['columns'],'gross_amount':'المطلوب','copay_amount':'إجمالي التحمل'}}
                policy=AuditPolicy('SYN-policy','LYD','HALF_UP',2,('exclusion','tariff','cap','copay'),Decimal('20'),None,Decimal('0'),Decimal('75' if i==1 else '100'),True,
                    amount_basis='NET',reference_fingerprint=reference_fingerprint(reference_check))
                policy_fp=fingerprint_policy(policy)
                calc=evaluate_net(claim,policy,reference_check)
                verification={'status':'verified'}
                if i==3:findings[0]['status']='review'
            db.add(AuditSnapshot(organization_id=claim.organization_id,claim_id=claim.id,policy_id='SYN-policy',
                policy_fingerprint=policy_fp,calculation=calc,findings={'agents':findings,'reference_check':reference_check},verification=verification,created_by=api.ids['root']))
        db.commit()


def test_arabic_mapping_imports_net_and_preserves_original_bytes(workspace):
    api=workspace;batch,doc,raw=imported(api)
    with Session(api.engine) as db:
        claims=list(db.scalars(select(ClaimEntry).where(ClaimEntry.batch_id==batch)))
        assert all(c.claimed_amount==Decimal('80') and c.service_date==date(2026,6,8) for c in claims)
        assert all(c.source_values['المطلوب']=='100' and c.source_values['الصافي']=='80' for c in claims)
        source=db.get(SourceFile,db.get(IngestionDocument,doc).source_id)
        assert source.sha256==sha256(raw).hexdigest()
        assert store().get(source.storage_key,source.sha256)==raw
    assert api.client.put('/api/v1/claims/batches/'+batch+'/mapping',json=mapping(),headers=api.headers['root']).status_code==409


def test_gross_column_is_rejected_and_mapping_cannot_cross_tenants(workspace):
    api=workspace;batch,_=upload_batch(api,arabic_workbook());payload=mapping();payload['columns']['claimed_amount']='المطلوب'
    assert api.client.put('/api/v1/claims/batches/'+batch+'/mapping',json=payload,headers=api.headers['root']).status_code==409
    assert api.client.get('/api/v1/claims/batches/'+batch+'/columns',headers=api.headers['other']).status_code==403


def test_four_files_are_disjoint_complete_and_idempotent(workspace):
    api=workspace;batch,doc,raw=imported(api);snapshots(api,batch)
    response=api.client.post('/api/v1/result-files/batches/'+batch,headers=api.headers['root'])
    assert response.status_code==201,response.text
    result=response.json();assert result['summary']['counts']=={'accepted':1,'rejected':1,'held':1,'review':1}
    assert result['summary']['total_lines']==4 and result['financially_final'] is False
    again=api.client.post('/api/v1/result-files/batches/'+batch,headers=api.headers['root'])
    assert again.json()['id']==result['id']
    download=api.client.get('/api/v1/result-files/'+result['id']+'/download',headers=api.headers['root'])
    assert download.status_code==200,download.text
    all_ids=[]
    with ZipFile(BytesIO(download.content)) as archive:
        assert set(archive.namelist())=={f'{key}.xlsx' for key in results.BUCKETS}
        for filename in archive.namelist():
            book=load_workbook(BytesIO(archive.read(filename)),read_only=False)
            rows=list(book.active.values);assert len(rows)==2
            assert book.active.sheet_view.rightToLeft
            all_ids.append(rows[1][0]);assert rows[1][1]=='SYN-SAME-INVOICE'
            assert all(cell!='SYN-BENEFICIARY' for row in rows for cell in row)
    assert len(all_ids)==len(set(all_ids))==4
    assert api.client.get('/api/v1/result-files/'+result['id']+'/lines',headers=api.headers['other']).status_code==403
    with Session(api.engine) as db:
        source=db.get(SourceFile,db.get(IngestionDocument,doc).source_id)
        assert store().get(source.storage_key,source.sha256)==raw


def test_no_audit_snapshot_never_becomes_an_accepted_result(workspace):
    api=workspace;batch,_,_=imported(api)
    from tests.workspace_v2.invoice_fixtures import attach_invoices
    attach_invoices(api,batch)
    response=api.client.post('/api/v1/result-files/batches/'+batch,headers=api.headers['root'])
    assert response.status_code==201,response.text
    assert response.json()['summary']['counts']=={'accepted':0,'rejected':0,'held':4,'review':0}
    rows=api.client.get('/api/v1/result-files/'+response.json()['id']+'/lines',headers=api.headers['root']).json()['items']
    assert all(row['audited'] is False and row['reason']=='not_audited' for row in rows)


def test_partition_rejects_missing_duplicate_and_unknown_lines():
    row={'line_id':'SYN-1','currency':'LYD','claimed_net':'80','bucket':'accepted'}
    with pytest.raises(ReportBlocked):results.partition([row,row],['SYN-1'])
    with pytest.raises(ReportBlocked):results.partition([row],['SYN-1','SYN-2'])
    with pytest.raises(ReportBlocked):results.partition([{**row,'bucket':'unknown'}],['SYN-1'])


def test_outputs_require_completed_import_and_do_not_modify_batch(workspace):
    api=workspace;batch,_=upload_batch(api,arabic_workbook())
    assert api.client.post('/api/v1/result-files/batches/'+batch,headers=api.headers['root']).status_code==409
    with Session(api.engine) as db:assert db.get(ClaimBatch,batch).imported_count==0
