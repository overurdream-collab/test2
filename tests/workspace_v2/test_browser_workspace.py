"""Browser -> real TestClient API -> isolated SQLite, with no external traffic."""
from pathlib import Path
import importlib.util
import shutil
import time
from urllib.parse import urlsplit
import pytest
from tests.api.conftest import api_factory, synthetic_password_hash, MFA_SECRET
from tests.features.conftest import features_api
from tests.workspace_v2.test_access_and_registry import workspace
from mcias.app import app, get_rate_limit_db
from mcias.mfa import generate_totp


@pytest.mark.skipif(not importlib.util.find_spec('playwright'),reason='Playwright browser dependency is unavailable')
def test_five_sections_reference_tables_and_permissions_in_browser(workspace,monkeypatch,tmp_path):
    from playwright.sync_api import sync_playwright
    from sqlalchemy.orm import Session
    from mcias.models import User
    from mcias.seed.reference_inputs import seed_reference_inputs
    api=workspace;monkeypatch.setenv('MCIAS_ENVIRONMENT','development');monkeypatch.setenv('MCIAS_LAPTOP_TEST_MODE','synthetic_localhost')
    with Session(api.engine) as db:
        seed_reference_inputs(db,db.get(User,api.ids['root']));db.commit()
    app.dependency_overrides[get_rate_limit_db]=api.database
    with sync_playwright() as p:
        executable=shutil.which('chromium')
        if not executable and not Path(p.chromium.executable_path).exists():pytest.skip('Chromium binary is unavailable')
        browser=p.chromium.launch(executable_path=executable,headless=True,args=['--no-sandbox'])
        page=browser.new_page(viewport={'width':1440,'height':1000},accept_downloads=True)
        errors=[];responses=[]
        page.on('pageerror',lambda error:errors.append(str(error)))
        def local_request(route):
            request=route.request;url=urlsplit(request.url);path=url.path+('?' + url.query if url.query else '')
            headers={k:v for k,v in request.headers.items() if k.lower() in ('authorization','content-type','accept')}
            response=api.client.request(request.method,path,content=request.post_data_buffer,headers=headers)
            responses.append((request.method,url.path,response.status_code))
            route.fulfill(status=response.status_code,headers={k:v for k,v in response.headers.items() if k.lower() not in ('content-length','content-encoding')},body=response.content)
        page.route('**/*',local_request)
        page.goto('http://mcias.test/login')
        page.locator('[name=username]').fill(api.root_username)
        page.locator('[name=password]').fill(api.password)
        page.locator('[name=totp_code]').fill(generate_totp(MFA_SECRET,int(time.time())//30))
        page.get_by_role('button',name='تسجيل الدخول بأمان',exact=True).click()
        page.wait_for_selector('#sidebar nav .nav-heading')
        assert page.locator('#sidebar nav .nav-heading').all_text_contents()==['الرئيسية','المدخلات','التدقيق والمراجعة','التقارير','الإدارة']
        assert not any('استيعاب' in name for name in page.locator('#sidebar .nav-link').all_text_contents())
        page.locator('#sidebar summary',has_text='البيانات والعقود').click()
        page.locator('#sidebar a[href="/institution/inputs/contracts"]').click()
        page.wait_for_selector('.table-scroll tbody tr')
        assert page.get_by_role('button',name='إضافة عقد جديد',exact=True).is_visible()
        assert page.get_by_role('button',name='تعديل بند',exact=True).is_disabled()
        assert page.get_by_role('button',name='إضافة بند',exact=True).is_disabled()
        assert page.get_by_role('button',name='استيراد ملف',exact=True).count()==0
        page.get_by_role('button',name='إضافة عقد جديد',exact=True).click()
        page.locator('[name=intake-title]').fill('SYN عقد المتصفح')
        page.locator('[name=intake-record_code]').fill('SYN-BROWSER-CONTRACT')
        page.locator('[name=intake-provider_code]').fill('SYN-BROWSER-HOSPITAL')
        page.locator('[name=intake-contract_number]').fill('SYN-BROWSER-2026')
        page.get_by_role('button',name='حفظ السجل',exact=True).click()
        page.get_by_role('button',name='SYN عقد المتصفح',exact=True).wait_for()
        row=page.locator('tr',has=page.get_by_role('button',name='SYN عقد المتصفح',exact=True))
        row.locator('input[type=radio]').check()
        page.get_by_role('button',name='تعديل بند',exact=True).click()
        page.locator('[name=intake-title]').fill('SYN عقد المتصفح المعدل')
        page.locator('[name=intake-edit-reason]').fill('SYN browser atomic replacement')
        page.get_by_role('button',name='حفظ السجل',exact=True).click()
        page.get_by_role('button',name='SYN عقد المتصفح المعدل',exact=True).wait_for()
        assert any(path.endswith('/replace') and status==201 for _,path,status in responses)
        # Source import, explicit Arabic mapping, four output files: actual APIs.
        from tests.workspace_v2.test_mapping_and_results import arabic_workbook
        page.locator('#sidebar a[href="/institution/audit-workspace"]').click()
        page.locator('[name=claims-file]').set_input_files({'name':'SYN-browser-claims.xlsx',
            'mimeType':'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet','buffer':arabic_workbook()})
        page.get_by_role('button',name='رفع ومعاينة',exact=True).click()
        page.locator('[name=map-currency]').fill('LYD')
        page.locator('[name=map-date]').select_option('DMY')
        assert page.locator('[name=map-claimed_amount]').input_value()=='الصافي'
        page.get_by_role('button',name='تأكيد الربط والمعاينة',exact=True).click()
        page.get_by_role('button',name='معاينة واستكمال الاستيراد',exact=True).click()
        page.get_by_role('button',name='تأكيد استيراد الصفوف الصالحة',exact=True).click()
        page.get_by_role('button',name='تجهيز ملفات النتائج الأربعة',exact=True).click()
        page.get_by_role('button',name='تنزيل الملفات الأربعة',exact=True).wait_for()
        with page.expect_download() as event:
            page.get_by_role('button',name='تنزيل الملفات الأربعة',exact=True).click()
        output=tmp_path/'results.zip';event.value.save_as(output)
        from zipfile import ZipFile
        with ZipFile(output) as archive:
            assert set(archive.namelist())=={'accepted.xlsx','rejected.xlsx','held.xlsx','review.xlsx'}
            from openpyxl import load_workbook
            from io import BytesIO
            held=list(load_workbook(BytesIO(archive.read('held.xlsx')),read_only=True).active.values)
            assert len(held)==5
            assert all('نقص المرفقات' in row[9] for row in held[1:])
            for name in ('accepted.xlsx','rejected.xlsx','review.xlsx'):
                assert load_workbook(BytesIO(archive.read(name)),read_only=True).active.max_row==1
        assert any(path.endswith('/resume') and status==200 for _,path,status in responses)
        # A second, linked workbook exercises the actual category/service/network
        # resolver rather than manufactured financial snapshots.
        from mcias.seed.linked_references import sample_workbook
        page.locator('#sidebar a[href="/institution/audit-workspace"]').click()
        page.locator('[name=claims-file]').set_input_files({'name':'SYN-linked-reference-claims.xlsx',
            'mimeType':'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet','buffer':sample_workbook()})
        page.get_by_role('button',name='رفع ومعاينة',exact=True).click()
        page.locator('[name=map-currency]').fill('LYD')
        page.locator('[name=map-date]').select_option('DMY')
        assert page.locator('[name=map-beneficiary_name]').input_value()=='اسم المستفيد'
        assert page.locator('[name=map-copay_amount]').input_value()=='إجمالي التحمل'
        page.get_by_role('button',name='تأكيد الربط والمعاينة',exact=True).click()
        page.get_by_role('button',name='معاينة واستكمال الاستيراد',exact=True).click()
        page.get_by_role('button',name='تأكيد استيراد الصفوف الصالحة',exact=True).click()
        page.get_by_text('مطابق: 0 · اختلاف: 0 · ناقص: 4',exact=True).wait_for()
        assert page.get_by_role('button',name='بدء التدقيق',exact=True).is_visible()
        page.get_by_role('button',name='تفاصيل الفحص',exact=True).first.click()
        assert page.locator('#drawer').get_by_text('نقص المرفقات — فاتورة الخدمة غير متوفرة',exact=True).is_visible()
        page.keyboard.press('Escape')
        from tests.workspace_v2.invoice_fixtures import attach_invoices
        batch_id=page.locator('[name=audit-batch]').input_value()
        attach_invoices(api,batch_id)
        page.get_by_role('button',name='فحص / إعادة فحص المدخلات',exact=True).click()
        page.get_by_text('مطابق: 4 · اختلاف: 0 · ناقص: 0',exact=True).wait_for()
        assert page.get_by_role('button',name='تفاصيل الفحص',exact=True).count()==4
        # A matching invoice/reference check is insufficient for financial
        # execution: no signed policy is enrolled in this browser fixture.
        page.get_by_role('button',name='بدء التدقيق',exact=True).click()
        page.get_by_role('heading',name='نتيجة دورة التدقيق',exact=True).wait_for()
        run=api.client.get(f'/api/v1/audits/batches/{batch_id}/run',headers=api.headers['root'])
        assert run.status_code==200,run.text
        assert run.json()['counts']=={'accepted':0,'rejected':0,'held':4,'review':0}
        assert all(item['reason']=='missing_or_ambiguous_signed_policy' and item['snapshot_id'] is None for item in run.json()['rows'])
        assert any(method=='POST' and path==f'/api/v1/audits/batches/{batch_id}/run' and status==201 for method,path,status in responses)
        evidence_dir=Path(__import__('os').getenv('MCIAS_BROWSER_ARTIFACT_DIR',str(tmp_path/'browser-evidence')))
        evidence_dir.mkdir(parents=True,exist_ok=True)
        page.screenshot(path=str(evidence_dir/'MCIAS_invoice_cycle_unsigned_ar.png'),full_page=True)
        # A non-admin stays frozen until the technical super-admin explicitly acts.
        page.locator('#sidebar a[href="/institution/administration"]').click()
        with Session(api.engine) as db:username=db.get(User,api.ids['admin']).username
        user_row=page.locator('tr',has_text=username)
        user_row.get_by_role('button',name='الصلاحيات',exact=True).click()
        page.locator('input[value="registry.inputs.contracts.read"]').check()
        page.locator('[name=access-reason]').fill('SYN external browser administrative decision')
        page.get_by_role('button',name='حفظ الصلاحيات',exact=True).click()
        page.locator('#drawer').wait_for(state='hidden')
        assert api.client.get('/api/v1/registry/inputs.contracts/records',headers=api.headers['admin']).status_code==200
        from tests.workspace_v2.test_access_and_registry import record
        assert api.client.post('/api/v1/registry/inputs.contracts/records',json=record('SYN-DENIED-CREATE'),headers=api.headers['admin']).status_code==403
        # Return to the reference table for the retained visual evidence.
        page.locator('#sidebar summary',has_text='البيانات والعقود').click()
        page.locator('#sidebar a[href="/institution/inputs/contracts"]').click()
        page.wait_for_selector('.table-scroll tbody tr')
        # Runtime colors/direction and mobile layout use the same real pages.
        assert page.locator('html').get_attribute('dir')=='rtl'
        page.set_viewport_size({'width':390,'height':844})
        assert page.locator('body').evaluate('(body)=>body.scrollWidth<=window.innerWidth')
        import os
        artifact=Path(os.getenv('MCIAS_BROWSER_ARTIFACT_DIR',str(tmp_path/'browser-evidence')))
        artifact.mkdir(parents=True,exist_ok=True)
        page.screenshot(path=str(artifact/'MCIAS_inputs_ar_mobile.png'),full_page=False)
        page.set_viewport_size({'width':1440,'height':1000})
        page.screenshot(path=str(artifact/'MCIAS_inputs_ar_desktop.png'),full_page=False)
        assert not errors,errors
        assert not any(status>=500 for _,_,status in responses),responses
        browser.close()
