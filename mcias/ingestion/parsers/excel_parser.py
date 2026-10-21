"""Read-only XLSX parser; never evaluates formulas or executes macros."""
import io
from openpyxl import load_workbook
from mcias.ingestion.schemas import Extraction, ExtractedRecord, Locator
from mcias.ingestion.quarantine import item

def parse(data: bytes) -> Extraction:
    book = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
    result = Extraction(format="xlsx")
    try:
        if len(book.worksheets) > 50:
            raise ValueError("worksheet limit exceeded")
        for sheet_index, sheet in enumerate(book.worksheets, start=1):
            inventory = {"index": sheet_index, "name": sheet.title, "headers": [], "data_rows": 0}
            result.worksheet_inventory.append(inventory)
            iterator = sheet.iter_rows(values_only=False)
            header = next(iterator, None)
            if not header:
                result.units_total += 1
                result.quarantined.append(item("empty worksheet", "sheet", sheet_index))
                continue
            if len(header) > 256:
                raise ValueError("worksheet column limit exceeded")
            names = [str(v.value).strip() if v.value is not None else "" for v in header]
            inventory["headers"] = names
            if not all(names) or len(names) != len(set(names)):
                result.units_total += 1
                result.quarantined.append(item("invalid worksheet header", "sheet", sheet_index))
                continue
            for row_index, values in enumerate(iterator, start=2):
                inventory["data_rows"] += 1
                result.units_total += 1
                if result.units_total > 100_000:
                    raise ValueError("spreadsheet row limit exceeded")
                if any(cell is not None and getattr(cell, "data_type", None) == "f" for cell in values):
                    result.quarantined.append(item("formula cell", "row", row_index))
                    continue
                result.records.append(ExtractedRecord(Locator("sheet", sheet_index, row_index, sheet_name=sheet.title), {key: "" if value is None else str(value) for key, value in zip(names, (cell.value for cell in values))}))
                result.units_processed += 1
    finally:
        book.close()
    return result
