"""G7 Excel/PDF rendering from verified, frozen snapshots. Never export source PHI by default."""
from io import BytesIO
from html import escape
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from mcias.reporting_g7 import verify_snapshot, ReportBlocked

def _text_cell(sheet, row, column, value):
    # Always store source values as literal strings; Excel must not execute "=..."
    # or coerce long identifiers to floating point.
    value = "" if value is None else str(value)
    if len(value) > 32767:
        raise ReportBlocked("source cell exceeds Excel text length; no silent truncation")
    cell = sheet.cell(row, column, value)
    cell.data_type = "s"
    cell.number_format = "@"
    return cell

def render_xlsx(report, *, include_originals=False, language="en"):
    if language not in {"ar", "en"}:
        raise ReportBlocked("unsupported report language")
    verify_snapshot(report)
    wb = Workbook()
    ws = wb.active
    ws.title = "الملخص التنفيذي" if language == "ar" else "Executive Summary"
    ws.sheet_view.rightToLeft = language == "ar"
    headings = [
        ["MCIAS | Frozen Financial Report"],
        ["Report ID", report.id], ["Version", report.version],
        ["Frozen SHA-256", report.fingerprint],
        ["Batch", report.batch_id], ["Source SHA-256", report.dataset["source_sha256"]],
        ["Total claims", report.summary["total_claims"]], [],
        ["Currency", "Claimed", "Approved claimed", "Held claimed", "Payable",
         "Deductions", "Approved count", "Held count"],
    ]
    if language == "ar":
        headings = [
            ["MCIAS | تقرير مراجعة المطالبات"],
            ["معرّف التقرير", report.id], ["رقم النسخة", report.version],
            ["بصمة النسخة SHA-256", report.fingerprint],
            ["الدفعة", report.batch_id], ["بصمة المصدر SHA-256", report.dataset["source_sha256"]],
            ["إجمالي المطالبات", report.summary["total_claims"]], [],
            ["العملة", "المطالب به", "المطالبات المعتمدة", "المطالبات المعلقة",
             "صافي المستحق", "الخصومات", "عدد المعتمدة", "عدد المعلقة"],
        ]
    for row in headings:
        ws.append(row)
    for currency, item in sorted(report.summary["by_currency"].items()):
        ws.append([currency, item["claimed"], item["approved_claimed"],
                   item["held_claimed"], item["payable"], item["deductions"],
                   item["approved_count"], item["held_count"]])
    ws.freeze_panes = "A10"
    ws.column_dimensions["A"].width = 29
    for col in "BCDEFGH":
        ws.column_dimensions[col].width = 23
    for cell in ws[9]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="243B53")
        cell.alignment = Alignment(horizontal="center")
    ws.sheet_view.showGridLines = False
    for row in ws:
        for cell in row:
            cell.alignment = Alignment(horizontal="right" if language == "ar" else "left")
    for cell in ws[9]:
        cell.alignment = Alignment(horizontal="center")
    if include_originals:
        raw = wb.create_sheet("القيم الأصلية" if language == "ar" else "Original Source Values")
        raw.sheet_view.rightToLeft = language == "ar"
        fields = sorted({k for item in report.dataset["rows"] for k in item["source_values"]})
        headers = (["معرّف المطالبة", "معرّف سجل المصدر", "موضع المصدر"]
                   if language == "ar" else ["Claim ID", "Source Record ID", "Source Locator"]) + fields
        for col, heading in enumerate(headers, start=1):
            _text_cell(raw, 1, col, heading)
            cell = raw.cell(1, col)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="243B53")
        import json
        for index, item in enumerate(report.dataset["rows"], start=2):
            values = [item["claim_id"], item["source_record_id"],
                      json.dumps(item["source_locator"], ensure_ascii=False, sort_keys=True)]
            values += [item["source_values"].get(field) for field in fields]
            for col, value in enumerate(values, start=1):
                _text_cell(raw, index, col, value)
        raw.freeze_panes = "D2"
        raw.auto_filter.ref = raw.dimensions
        for col in raw.columns:
            raw.column_dimensions[col[0].column_letter].width = min(38, max(16, len(str(col[0].value)) + 2))
    output = BytesIO()
    wb.save(output)
    return output.getvalue()

def render_pdf(report):
    verify_snapshot(report)
    output = BytesIO()
    doc = SimpleDocTemplate(output, pagesize=A4, title="MCIAS Frozen Report",
                            leftMargin=36, rightMargin=36, topMargin=44, bottomMargin=40)
    style = getSampleStyleSheet()
    story = [Paragraph("MCIAS | Executive Financial Report", style["Title"]),
             Spacer(1, 14),
             Paragraph("Frozen version: " + str(report.version), style["Normal"]),
             Paragraph("Batch: " + escape(report.batch_id), style["Normal"]),
             Paragraph("SHA-256: " + escape(report.fingerprint), style["Normal"]),
             Paragraph("Total claims: " + str(report.summary["total_claims"]), style["Normal"]),
             Spacer(1, 16)]
    rows = [["Currency", "Claimed", "Approved", "Held", "Payable", "Deductions"]]
    for currency, item in sorted(report.summary["by_currency"].items()):
        rows.append([currency, item["claimed"], item["approved_claimed"],
                     item["held_claimed"], item["payable"], item["deductions"]])
    table = Table(rows, repeatRows=1, hAlign="LEFT", colWidths=[64, 91, 91, 91, 91, 91])
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#243B53")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 9),
        ("TOPPADDING", (0, 0), (-1, -1), 9),
        ("LINEBELOW", (0, 0), (-1, 0), 0.5, colors.HexColor("#D7E2EF")),
    ]))
    story += [table, Spacer(1, 18),
              Paragraph("Held items are never treated as approved payable amounts.", style["Normal"]),
              Paragraph("Source/patient identifiers are intentionally absent from this PDF.", style["Normal"])]
    doc.build(story)
    return output.getvalue()

def render_result_partition(dataset, bucket):
    """One sheet, one category, one row per source line; no duplicate all-claims file."""
    from mcias.services.result_service import BUCKETS, LABELS, partition
    if bucket not in BUCKETS:raise ReportBlocked('Unsupported output category')
    partition(dataset['rows'],[row['line_id'] for row in dataset['rows']])
    wb=Workbook();sheet=wb.active;sheet.title=LABELS[bucket];sheet.sheet_view.rightToLeft=True
    columns=[('line_id','معرف البند'),('invoice_number','رقم الفاتورة'),('provider_code','مقدم الخدمة'),
        ('service_code','الخدمة'),('service_date','تاريخ الخدمة'),('currency','العملة'),
        ('claimed_net','الصافي الأصلي'),('proposed_net','الصافي المقترح'),
        ('proposed_difference','الفرق المقترح'),('reason','سبب النتيجة'),('audited','تم تدقيقه'),
        ('audit_snapshot_id','مرجع لقطة التدقيق'),('source_record_id','سجل المصدر')]
    for col,(_,title) in enumerate(columns,1):
        cell=_text_cell(sheet,1,col,title);cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='243B53')
    for row_index,row in enumerate((r for r in dataset['rows'] if r['bucket']==bucket),2):
        for col,(key,_) in enumerate(columns,1):_text_cell(sheet,row_index,col,row.get('reason_ar',row.get(key)) if key=='reason' else row.get(key))
    sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
    for col in sheet.columns:sheet.column_dimensions[col[0].column_letter].width=24
    sheet.sheet_properties.pageSetUpPr.fitToPage=True
    output=BytesIO();wb.save(output);return output.getvalue()
