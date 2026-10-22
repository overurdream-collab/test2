"""One synthetic claim through real HTTP services and optional browser controls.

Only account/role/MFA setup writes directly to the disposable database. Documents,
reference approval, claims, G5, Shared Audit, G6 and G7 use the application's HTTP
entry points. Independent database sessions only verify committed evidence.
"""
import argparse
import base64
from contextlib import ExitStack
import csv
from decimal import Decimal
import hashlib
import io
import json
import os
from pathlib import Path
import time
from urllib.parse import urlsplit
from uuid import uuid4

from cryptography.hazmat.primitives.ciphers.aead import AESGCM
from openpyxl import Workbook, load_workbook
from sqlalchemy import create_engine, select, text
from sqlalchemy.engine import make_url
from sqlalchemy.orm import Session

from mcias.audit_agents import project_audit_findings
from mcias.mfa import generate_totp
from mcias.models import (AuditLog, MFAFactor, Organization, Permission, Role,
                          RolePermission, SourceFile, User, UserRole)
from mcias.models_audit import AuditSnapshot
from mcias.models_claims import ClaimEntry
from mcias.models_governance_g6 import GovernanceEvent
from mcias.models_ingestion import ExtractedSourceRecord, IngestionDocument, IngestionJob
from mcias.models_reporting_g7 import ReportSnapshot
from mcias.security import hash_password

ROOT = Path(__file__).resolve().parents[1]
PASSWORD = 'synthetic_workflow_password_only'
SHEET_TYPES = {'1': 'CLAIMS', '2': 'PAYMENT', '3': 'FAMILY_SUMMARY'}
FAMILIES = ('beneficiary', 'contract', 'tariff', 'policy', 'caps', 'history')


def _csv(rows):
    output = io.StringIO()
    writer = csv.DictWriter(output, fieldnames=list(rows[0]))
    writer.writeheader()
    writer.writerows(rows)
    return output.getvalue().encode('utf-8')


def _workbook(claim_reference):
    row = {'claim_reference': claim_reference, 'beneficiary_token': 'SYN-BEN',
           'provider_code': 'SYN-PROVIDER', 'service_code': 'SYN-SERVICE',
           'service_date': '2026-06-01', 'claimed_amount': '120', 'currency': 'USD',
           'quantity': '1', 'unit_price': '120', 'unit': 'UNIT', 'copay_amount': '24',
           'diagnosis': 'SYNTHETIC-PRIVATE-DIAGNOSIS'}
    book = Workbook()
    for index, name in enumerate(('CLAIMS', 'PAYMENT', 'FAMILY_SUMMARY')):
        sheet = book.active if index == 0 else book.create_sheet()
        sheet.title = name
        sheet.append(list(row))
        sheet.append([value if key != 'claim_reference' or index == 0 else 'SYN-' + name
                      for key, value in row.items()])
    buffer = io.BytesIO()
    book.save(buffer)
    book.close()
    return buffer.getvalue()


def _reference_facts():
    dates = {'effective_from': '2026-01-01', 'effective_until': '2026-07-01'}
    facts = {
        'beneficiary': {'beneficiary_token': 'SYN-BEN', 'employee_token': 'SYN-EMP',
            'relationship': 'self', 'category_code': 'SYN-CAT', 'eligible': True,
            'covered_services': ['SYN-SERVICE']},
        'contract': {'provider_code': 'SYN-PROVIDER', 'contract_key': 'SYN-CONTRACT',
            'covered_services': ['SYN-SERVICE'], 'excluded_services': []},
        'tariff': {'provider_code': 'SYN-PROVIDER', 'service_code': 'SYN-SERVICE',
            'category_code': 'SYN-CAT', 'currency': 'USD', 'unit': 'UNIT',
            'unit_price': '100', 'price_basis': 'unit_price', 'price_comparison': 'exact'},
        'policy': {'service_code': 'SYN-SERVICE', 'category_code': 'SYN-CAT',
            'currency': 'USD', 'copay_percent': '20', 'requires_execution_evidence': True,
            'minimum_dispense_interval_days': None, 'copay_basis': 'claimed_amount',
            'rounding': 'HALF_UP', 'precision': 2, 'duplicate_paid_claims_prohibited': True,
            'early_dispense_days': 2},
        'caps': {'beneficiary_token': 'SYN-BEN', 'category_code': 'SYN-CAT',
            'service_code': 'SYN-SERVICE', 'currency': 'USD', 'limit_amount': '1000',
            'consumed_amount': '0', 'period_from': '2026-01-01', 'period_until': '2026-07-01',
            'consumption_as_of': '2026-06-01', 'consumed_excludes_current_claim': True,
            'consumption_basis': 'claimed_amount'},
        'history': {'beneficiary_token': 'SYN-BEN', 'provider_code': 'SYN-PROVIDER',
            'service_code': 'SYN-OTHER', 'service_date': '2026-05-01', 'claim_reference': 'SYN-OLD',
            'paid': True, 'quantity': '1', 'unit': 'UNIT', 'amount': '100', 'currency': 'USD'}}
    return {family: {**value, **dates} for family, value in facts.items()}


def _accounts(engine, unique):
    """Bootstrap only genuine synthetic security identities; never domain rows."""
    operator_permissions = ['profile.read', 'claims.read', 'claims.import', 'ingestion.read',
        'ingestion.upload', 'reports.read', 'reports.create', 'reports.export',
        'reports.details.read', 'workspace.super_admin', 'reference.read', 'reference.write',
        'reference.approve', 'governance.review']
    ids = {}
    secret = os.urandom(20)
    with Session(engine) as db:
        organization = Organization(name='SYNTHETIC WORKFLOW ' + unique)
        foreign = Organization(name='SYNTHETIC FOREIGN ' + unique)
        db.add_all([organization, foreign])
        db.flush()
        for name, codes in (('operator', operator_permissions),
                            ('reviewer', ['profile.read', 'claims.read', 'governance.review']),
                            ('foreign', ['profile.read', 'claims.read', 'reports.read', 'ingestion.read'])):
            role = Role(code='workflow_' + name + '_' + unique)
            db.add(role)
            db.flush()
            for code in codes:
                permission = db.scalar(select(Permission).where(Permission.code == code))
                if permission is None:
                    permission = Permission(code=code)
                    db.add(permission)
                    db.flush()
                db.add(RolePermission(role_id=role.id, permission_id=permission.id))
            user = User(organization_id=foreign.id if name == 'foreign' else organization.id,
                        username='workflow-' + name + '-' + unique, password_hash=hash_password(PASSWORD))
            db.add(user)
            db.flush()
            db.add(UserRole(user_id=user.id, role_id=role.id))
            ids[name] = {'id': user.id, 'username': user.username, 'organization_id': user.organization_id}
        administrator = db.scalar(select(Role).where(Role.code == 'super_admin'))
        if administrator is None:
            administrator = Role(code='super_admin')
            db.add(administrator)
            db.flush()
        # Approval trigger 0019 requires the permission on this canonical role,
        # rather than on a different role belonging to the same account.
        workspace_permission = db.scalar(select(Permission).where(
            Permission.code == 'workspace.super_admin'))
        canonical_grant = db.scalar(select(RolePermission).where(
            RolePermission.role_id == administrator.id,
            RolePermission.permission_id == workspace_permission.id))
        if canonical_grant is None:
            db.add(RolePermission(role_id=administrator.id, permission_id=workspace_permission.id))
        db.add(UserRole(user_id=ids['operator']['id'], role_id=administrator.id))
        nonce = os.urandom(12)
        encrypted = AESGCM(bytes.fromhex(os.environ['MCIAS_MFA_ENCRYPTION_KEY'])).encrypt(
            nonce, secret, ids['operator']['id'].encode())
        db.add(MFAFactor(user_id=ids['operator']['id'], encrypted_secret=nonce + encrypted, enabled=True))
        db.commit()
    return ids, secret


def _private_manifest(directory, ids, secret, base_url):
    path = directory / 'integrated-workflow-credentials.json'
    descriptor = os.open(path, os.O_CREAT | os.O_TRUNC | os.O_WRONLY, 0o600)
    os.chmod(path, 0o600)
    with os.fdopen(descriptor, 'w', encoding='utf-8') as stream:
        json.dump({'synthetic_only': True, 'base_url': base_url,
                   'accounts': {name: {'username': account['username'], 'password': PASSWORD,
                       **({'totp_secret_base32': base64.b32encode(secret).decode('ascii')}
                          if name == 'operator' else {})} for name, account in ids.items()}},
                  stream, indent=2)


class HTTP:
    def __init__(self, client):
        self.client = client
        self.headers = {}
        self.statuses = []

    def request(self, method, path, *, headers=None, data=None, raw=None, expected=200):
        kwargs = {'headers': self.headers if headers is None else headers}
        if raw is not None:
            kwargs['content'] = raw
        elif data is not None:
            kwargs['json'] = data
        response = self.client.request(method, path, **kwargs)
        accepted = (expected,) if isinstance(expected, int) else tuple(expected)
        self.statuses.append({'method': method, 'path': path, 'status': response.status_code})
        assert response.status_code in accepted, (method, path, response.status_code, accepted)
        if 'json' in response.headers.get('content-type', ''):
            return response.json()
        return response.content

    def upload(self, filename, data):
        from urllib.parse import quote
        return self.request('POST', '/api/v1/ingestion/uploads?filename=' + quote(filename), raw=data)


def _browser_login(page, base_url, operator, secret):
    from playwright.sync_api import expect
    page.goto(base_url + '/login')
    page.get_by_label('اسم المستخدم', exact=True).fill(operator['username'])
    page.get_by_label('كلمة المرور', exact=True).fill(PASSWORD)
    page.get_by_label('رمز التحقق (إلزامي للإدارة)', exact=True).fill(generate_totp(secret, int(time.time()) // 30))
    with page.expect_response(lambda response: response.url.endswith('/api/v1/auth/login')) as result:
        page.get_by_role('button', name='تسجيل الدخول بأمان', exact=True).click()
    assert result.value.status == 200
    expect(page.locator('main h1')).to_be_visible()
    return result.value.json()['access_token']


def _navigate(page, path):
    from playwright.sync_api import expect
    page.evaluate('(path)=>{history.pushState({},"",path);dispatchEvent(new PopStateEvent("popstate"));}', path)
    expect(page.locator('main h1')).to_be_visible()


def _browser_operation(page, title, values, endpoint, *, confirmation=False, expected=200):
    form = page.locator('.form-panel').filter(has=page.get_by_role('heading', name=title, exact=True))
    for label, value in values.items():
        form.get_by_label(label, exact=True).fill(str(value))
    with page.expect_response(lambda response: endpoint in response.url and response.request.method == 'POST') as result:
        form.get_by_role('button', name='تنفيذ الطلب', exact=True).click()
        if confirmation:
            page.get_by_role('button', name='تأكيد الإجراء', exact=True).click()
    accepted = (expected,) if isinstance(expected, int) else tuple(expected)
    assert result.value.status in accepted, (endpoint, result.value.status, accepted)
    return result.value.json()


def _claim_import(api, page, original):
    if page is None:
        uploaded = api.upload('synthetic-workflow.xlsx', original)
    else:
        _navigate(page, '/admin/ingestion')
        page.get_by_label('المستند المصدر', exact=True).set_input_files({
            'name': 'synthetic-workflow.xlsx',
            'mimeType': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', 'buffer': original})
        with page.expect_response(lambda response: '/api/v1/ingestion/uploads?' in response.url) as response:
            page.get_by_role('button', name='رفع وتسجيل المصدر', exact=True).click()
        assert response.value.status == 200
        uploaded = response.value.json()
    assert uploaded['status'] == 'extracted'
    assert uploaded['sha256'] == hashlib.sha256(original).hexdigest()
    assert api.upload('synthetic-workflow.xlsx', original)['status'] == 'duplicate'
    api.request('GET', '/api/v1/ingestion/jobs/' + uploaded['job_id'] + '/records', expected=403)
    document_id = uploaded['document_id']
    if page is None:
        approval = api.request('POST', '/api/v1/claims/documents/' + document_id + '/source-scope',
                               data={'sheet_types': SHEET_TYPES}, expected=201)
        batch = api.request('POST', '/api/v1/claims/batches/from-document/' + document_id, expected=201)
        outcome = api.request('POST', '/api/v1/claims/batches/' + batch['id'] + '/resume')
    else:
        _navigate(page, '/institution/claims/batches')
        approval = _browser_operation(page, 'اعتماد نطاق أوراق المصدر',
            {'المستند': document_id, 'أنواع الأوراق بصيغة JSON': json.dumps(SHEET_TYPES)},
            '/source-scope', confirmation=True, expected=201)
        batch = _browser_operation(page, 'تجهيز دفعة من مصدر موثّق', {'المستند': document_id},
                                   '/batches/from-document/', expected=201)
        outcome = _browser_operation(page, 'استئناف استيراد الدفعة', {'الدفعة': batch['id']},
                                     '/resume?', confirmation=True)
    preview = api.request('GET', '/api/v1/claims/batches/' + batch['id'] + '/preview')
    assert preview['record_count'] == 1 and preview['excluded_count'] == 2
    assert len(preview['rows']) == 1 and preview['rows'][0]['valid'] is True
    preview_locator = preview['rows'][0]['locator']
    assert preview_locator['index'] == 1 and preview_locator['sheet_name'] == 'CLAIMS' and preview_locator['row'] == 2
    assert batch['record_count'] == 1
    assert outcome['status'] == 'imported' and outcome['imported_count'] == 1
    assert api.request('POST', '/api/v1/claims/batches/' + batch['id'] + '/resume')['imported_count'] == 1
    entries = api.request('GET', '/api/v1/claims/batches/' + batch['id'] + '/entries')['entries']
    assert len(entries) == 1
    locator = entries[0]['source_locator']
    assert locator['index'] == 1 and locator['sheet_name'] == 'CLAIMS' and locator['row'] == 2
    return uploaded, approval, batch, entries[0]


def _references(api, unique):
    result = []
    for family, facts in _reference_facts().items():
        original = _csv([{'record_key': 'SYN-ROW', 'payload': json.dumps(facts)}])
        uploaded = api.upload('SYN-' + family + '.csv', original)
        assert uploaded['status'] == 'extracted'
        item = api.request('POST', '/api/v1/reference/datasets', data={
            'document_id': uploaded['document_id'], 'family': family, 'dataset_key': 'SYN-' + family + '-' + unique,
            'version': 1, 'effective_from': '2026-01-01', 'effective_until': '2026-07-01',
            'coverage_from': '2026-01-01', 'coverage_until': '2026-07-01',
            'scope': {'provider_code': 'SYN-PROVIDER', 'all_beneficiaries': True, 'all_services': True},
            'expected_record_count': 1}, expected=201)
        api.request('POST', '/api/v1/reference/datasets/' + item['id'] + '/approve',
                    data={'reason_code': 'source_scope_complete'}, expected=(200, 201))
        durable = api.request('GET', '/api/v1/reference/datasets/' + item['id'])
        assert durable['status'] == 'approved' and durable['version'] == 1
        assert durable['source_sha256'] == uploaded['sha256'] == hashlib.sha256(original).hexdigest()
        result.append({'family': family, 'id': item['id'], 'version': 1,
                       'source_document_id': uploaded['document_id'], 'source_sha256': uploaded['sha256']})
    # The newer tariff is intentionally present but is not applicable in June.
    later = {**_reference_facts()['tariff'], 'effective_from': '2026-07-01',
             'effective_until': '2027-01-01', 'unit_price': '999'}
    uploaded = api.upload('SYN-tariff-v2.csv', _csv([{'record_key': 'SYN-ROW', 'payload': json.dumps(later)}]))
    newer = api.request('POST', '/api/v1/reference/datasets', data={
        'document_id': uploaded['document_id'], 'family': 'tariff', 'dataset_key': 'SYN-tariff-' + unique,
        'version': 2, 'effective_from': '2026-07-01', 'effective_until': '2027-01-01',
        'coverage_from': '2026-07-01', 'coverage_until': '2027-01-01',
        'scope': {'provider_code': 'SYN-PROVIDER', 'all_beneficiaries': True, 'all_services': True},
        'expected_record_count': 1}, expected=201)
    api.request('POST', '/api/v1/reference/datasets/' + newer['id'] + '/approve',
                data={'reason_code': 'source_scope_complete'}, expected=(200, 201))
    return result, newer['id']


def _dashboard(api, batch_id):
    report = api.request('POST', '/api/v1/reports/batches/' + batch_id + '/snapshots', expected=201)
    reread = api.request('GET', '/api/v1/reports/snapshots/' + report['id'] + '/dashboard')
    assert reread['fingerprint'] == report['fingerprint']
    summary = reread['summary']
    assert summary['total_claims'] == 1
    assert summary['audit']['outcome_counts']['VIOLATION_PROVEN'] == 1
    assert summary['audit']['completed_audit_count'] == 1
    assert summary['audit']['complete_reference_count'] == 6
    assert summary['audit']['source_evidence_verified_count'] == 1
    dashboard = summary['dashboard']
    assert set(dashboard['kpis']) == {'total_claimed', 'audited_claims', 'review_required', 'potential_exposure'}
    assert Decimal(dashboard['kpis']['total_claimed']['by_currency']['USD']['amount']) == Decimal('120')
    assert dashboard['kpis']['audited_claims']['count'] == 1
    assert dashboard['kpis']['review_required']['claim_count'] == 1
    assert dashboard['kpis']['review_required']['finding_count'] >= 1
    assert dashboard['kpis']['potential_exposure']['by_currency']['USD'] == {
        'status': 'UNQUANTIFIED', 'amount': None, 'unquantified_claim_count': 1}
    assert dashboard['classification_counts'] == {'NO_DETECTED_FINDINGS': 0,
        'HUMAN_REVIEW_REQUIRED': 1, 'MISSING_EVIDENCE': 0, 'UNABLE_TO_AUDIT': 0,
        'HUMAN_APPROVED_REJECTION': 0}
    assert dashboard['recommended_rejection_pending'] == {'available': False, 'claim_count': 0,
        'reason_code': 'APPROVED_REJECTION_POLICY_UNAVAILABLE'}
    assert dashboard['human_approved_rejected']['claim_count'] == 0
    assert summary['by_currency']['USD']['approved_count'] == 0
    assert Decimal(summary['by_currency']['USD']['payable']) == Decimal('0')
    exported = api.request('GET', '/api/v1/reports/snapshots/' + report['id'] + '/excel')
    assert exported.startswith(b'PK')
    book = load_workbook(io.BytesIO(exported), read_only=True)
    try:
        assert all('SYNTHETIC-PRIVATE-DIAGNOSIS' not in str(value)
                   for sheet in book for row in sheet.iter_rows(values_only=True) for value in row)
    finally:
        book.close()
    api.request('GET', '/api/v1/reports/snapshots/' + report['id'] + '/excel?include_originals=true', expected=403)
    return reread


def run(base_url, chromium=None, http_only=False):
    import httpx
    database_url = os.environ['MCIAS_E2E_DATABASE_URL']
    parsed = make_url(database_url)
    if parsed.get_backend_name() != 'postgresql' or not parsed.database or not parsed.database.endswith('_e2e'):
        raise ValueError('Dedicated PostgreSQL *_e2e database required')
    if os.environ.get('MCIAS_ENVIRONMENT') != 'development':
        raise ValueError('Explicit development environment required')
    if urlsplit(base_url).hostname not in {'localhost', '127.0.0.1', 'app'}:
        raise ValueError('Only the authorized local synthetic application is supported')
    output = ROOT / '.e2e-output'
    output.mkdir(parents=True, exist_ok=True)
    result = {'synthetic_data': True, 'actual_http_and_postgresql': True,
              'mode': 'http_only' if http_only else 'browser_and_http', 'status': 'RUNNING',
              'audit_execution': 'real POST /api/v1/audits/claims/{claim_id}/review',
              'institutional_acceptance': False, 'checks': [], 'page_errors': []}
    engine = create_engine(database_url)
    api = None
    try:
        with engine.connect() as connection:
            actual_version = connection.scalar(text('SHOW server_version'))
            assert int(connection.scalar(text('SHOW server_version_num'))) // 10000 == 16
        result['actual_postgresql_version'] = actual_version
        unique = uuid4().hex[:12]
        claim_reference = 'SYN-CLAIM-' + unique
        ids, secret = _accounts(engine, unique)
        _private_manifest(output, ids, secret, base_url)
        with ExitStack() as stack:
            client = stack.enter_context(httpx.Client(base_url=base_url.rstrip('/'), timeout=45,
                                                     trust_env=False))
            api = HTTP(client)
            result['http_statuses'] = api.statuses
            page = None
            if http_only:
                login = api.request('POST', '/api/v1/auth/login', headers={}, data={
                    'username': ids['operator']['username'], 'password': PASSWORD,
                    'totp_code': generate_totp(secret, int(time.time()) // 30)})
                token = login['access_token']
            else:
                from playwright.sync_api import sync_playwright, expect
                playwright = stack.enter_context(sync_playwright())
                browser = playwright.chromium.launch(executable_path=chromium,
                    args=['--no-sandbox'] if chromium else [])
                stack.callback(browser.close)
                page = browser.new_page(viewport={'width': 1440, 'height': 1000})
                page.on('pageerror', lambda error: result['page_errors'].append(str(error)))
                token = _browser_login(page, base_url, ids['operator'], secret)
            api.headers = {'Authorization': 'Bearer ' + token}
            result['checks'].append('real login with enrolled encrypted Super Admin MFA')
            original = _workbook(claim_reference)
            (output / 'integrated-workflow-claims.xlsx').write_bytes(original)
            uploaded, approval, batch, claim = _claim_import(api, page, original)
            claim_id = claim['id']
            result.update(organization_id=ids['operator']['organization_id'], claim_id=claim_id,
                claim_reference=claim_reference, source_document_id=uploaded['document_id'],
                source_sha256=uploaded['sha256'], source_scope_approval_id=approval['approval_id'],
                batch_id=batch['id'], source_locator=claim['source_locator'])
            result['checks'].append('real XLSX upload, SHA256, three-sheet approval; only CLAIMS imported and retry deduplicated')
            references, newer_tariff_id = _references(api, unique)
            result['references'] = references
            result['newer_inapplicable_tariff_id'] = newer_tariff_id
            execution = _csv([{'evidence_kind': 'service_execution', 'claim_reference': claim_reference,
                'beneficiary_token': 'SYN-BEN', 'provider_code': 'SYN-PROVIDER', 'service_code': 'SYN-SERVICE',
                'service_date': '2026-06-01', 'unit': 'UNIT', 'performed_quantity': '1', 'dose_status': 'clear'}])
            execution_doc = api.upload('SYN-execution.csv', execution)
            assert execution_doc['status'] == 'extracted'
            api.request('POST', '/api/v1/claims/entries/' + claim_id + '/attachments/' + execution_doc['document_id'], expected=201)
            result['execution_document_id'] = execution_doc['document_id']
            result['checks'].append('six source-backed reference datasets approved through HTTP, plus future tariff and execution evidence')
            review_path = '/api/v1/audits/claims/' + claim_id + '/review'
            if page is None:
                snapshot = api.request('POST', review_path, expected=(200, 201))
            else:
                _navigate(page, '/institution/financial-audit')
                snapshot = _browser_operation(page, 'تشغيل تدقيق المطالبة', {'المطالبة': claim_id},
                                               review_path, confirmation=True, expected=(200, 201))
            snapshot_id = snapshot['id']
            assert snapshot['claim_id'] == claim_id
            assert snapshot['findings']['audit_state'] == 'VIOLATION_PROVEN'
            assert snapshot['findings']['audit_state_schema_version'] == 2
            assert snapshot['findings']['coverage'] == {'required': 8, 'executed': 8, 'blocked': 0}
            assert snapshot['calculation']['status'] == 'held' and snapshot['calculation']['payable'] is None
            assert snapshot['calculation']['deductions'] == [] and snapshot['verification']['status'] == 'held'
            certificates = snapshot['findings']['reference_certificates']
            assert len(certificates) == 6 and all(item['approved'] and item['complete'] for item in certificates)
            tariff = next(item for item in certificates if item['family'] == 'tariff')
            assert tariff['reference_id'] == next(item['id'] for item in references if item['family'] == 'tariff')
            assert tariff['version'] == 1 and tariff['reference_id'] != newer_tariff_id
            assert api.request('POST', review_path, expected=(200, 201))['id'] == snapshot_id
            assert api.request('GET', '/api/v1/audits/' + snapshot_id)['findings'] == snapshot['findings']
            result.update(snapshot_id=snapshot_id, expected_audit_state='VIOLATION_PROVEN',
                          actual_audit_state=snapshot['findings']['audit_state'])
            result['checks'].append('real G5 HTTP execution proves dated tariff violation; repeated execution returns same held advisory snapshot')
            case = api.request('POST', '/api/v1/governance/cases/from-audit/' + snapshot_id, expected=201)
            reviewer = api.request('POST', '/api/v1/auth/login', headers={}, data={
                'username': ids['reviewer']['username'], 'password': PASSWORD})
            reviewer_headers = {'Authorization': 'Bearer ' + reviewer['access_token']}
            hold = api.request('POST', '/api/v1/governance/cases/' + case['id'] + '/events',
                headers=reviewer_headers, data={'kind': 'hold', 'payload': {'reason_code': 'SYNTHETIC_REVIEW_REQUIRED'}}, expected=201)
            api.request('POST', '/api/v1/governance/cases/' + case['id'] + '/events',
                headers=reviewer_headers, data={'kind': 'approve', 'payload': {}}, expected=403)
            history = api.request('GET', '/api/v1/governance/cases/' + case['id'])
            assert [event['kind'] for event in history['events']] == ['hold']
            result.update(governance_case_id=case['id'], human_review_event_id=hold['id'])
            result['checks'].append('separate real reviewer records G6 hold; financial approval remains unauthorized')
            report = _dashboard(api, batch['id'])
            result.update(report_id=report['id'], report_fingerprint=report['fingerprint'], dashboard=report['summary']['dashboard'])
            foreign = api.request('POST', '/api/v1/auth/login', headers={}, data={
                'username': ids['foreign']['username'], 'password': PASSWORD})
            foreign_headers = {'Authorization': 'Bearer ' + foreign['access_token']}
            api.request('GET', '/api/v1/workspace/claims/' + claim_id, headers=foreign_headers, expected=404)
            api.request('GET', '/api/v1/audits/' + snapshot_id, headers=foreign_headers, expected=404)
            api.request('GET', '/api/v1/reports/snapshots/' + report['id'], headers=foreign_headers, expected=404)
            result['checks'].append('real G7 four KPIs/five classifications, safe XLSX export, denied originals export and cross-tenant denial')
            with Session(engine) as independent:
                durable_claim = independent.get(ClaimEntry, claim_id)
                assert durable_claim.source_values['claim_reference'] == claim_reference
                assert durable_claim.organization_id == ids['operator']['organization_id']
                assert len(independent.scalars(select(ClaimEntry).where(ClaimEntry.batch_id == batch['id'])).all()) == 1
                source_records = independent.scalars(select(ExtractedSourceRecord).join(IngestionJob).where(
                    IngestionJob.document_id == uploaded['document_id'])).all()
                assert len(source_records) == 3
                assert {record.locator['sheet_name'] for record in source_records} == {'CLAIMS', 'PAYMENT', 'FAMILY_SUMMARY'}
                document = independent.get(IngestionDocument, uploaded['document_id'])
                assert independent.get(SourceFile, document.source_id).sha256 == hashlib.sha256(original).hexdigest()
                durable_snapshot = independent.get(AuditSnapshot, snapshot_id)
                assert durable_snapshot.findings['audit_state'] == 'VIOLATION_PROVEN'
                assert durable_snapshot.findings['source_evidence'] == snapshot['findings']['source_evidence']
                assert len(durable_snapshot.findings['reference_certificates']) == 6
                assert project_audit_findings(durable_snapshot.findings) == snapshot['findings']
                events = independent.scalars(select(AuditLog).where(AuditLog.action == 'claim_inputs_audited',
                    AuditLog.resource_id == claim_id)).all()
                assert len(events) == 1 and events[0].details['audit_snapshot_id'] == snapshot_id
                assert events[0].details['audit_state'] == 'VIOLATION_PROVEN'
                assert all(value not in json.dumps(events[0].details) for value in ('SYN-BEN', 'SYN-PROVIDER', claim_reference, 'SYNTHETIC-PRIVATE-DIAGNOSIS'))
                assert independent.get(GovernanceEvent, hold['id']).snapshot_id == snapshot_id
                durable_report = independent.get(ReportSnapshot, report['id'])
                assert durable_report.summary == report['summary']
                assert durable_report.dataset['rows'][0]['financial_decision_status'] == 'PENDING_HUMAN_REVIEW'
                result.update(audit_log_id=events[0].id,
                    source_record_id=durable_claim.extracted_record_id,
                    financial_decision_status=durable_report.dataset['rows'][0]['financial_decision_status'])
            result['checks'].append('independent PostgreSQL session rereads committed source, one shared audit event, G6 event and frozen G7 snapshot')
            if page is not None:
                _navigate(page, '/institution/dashboard')
                page.get_by_label('لقطة التقرير', exact=True).select_option(report['id'])
                expect(page.locator('.kpi')).to_have_count(4)
                for label in ('إجمالي قيمة المطالبات', 'المطالبات المدققة', 'تحتاج مراجعة بشرية', 'التعرض المالي المحتمل'):
                    expect(page.locator('.kpi-label').filter(has_text=label)).to_be_visible()
                expect(page.locator('main')).to_contain_text('غير متاح؛ لا سياسة رفض معتمدة')
                page.screenshot(path=str(output / 'integrated-workflow-dashboard.png'), full_page=True)
                assert not result['page_errors'], result['page_errors']
                result['checks'].append('real browser displays four contracted KPIs and five audit classifications without automatic rejection')
            result['http_statuses'] = api.statuses
            result['status'] = 'PASS'
    except Exception as error:
        result['status'] = 'FAIL'
        result['failure_type'] = type(error).__name__
        result['failure'] = str(error) if isinstance(error, AssertionError) else 'Runtime failure; inspect the private process log.'
        raise
    finally:
        if api is not None:
            result['http_statuses'] = api.statuses
        engine.dispose()
        (output / 'integrated-workflow-results.json').write_text(json.dumps(result, indent=2, ensure_ascii=False) + '\n', encoding='utf-8')
    print(json.dumps(result, indent=2, ensure_ascii=False))
    return result


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('--base-url', default='http://127.0.0.1:8002')
    parser.add_argument('--chromium', default=os.getenv('MCIAS_CHROMIUM'))
    parser.add_argument('--http-only', action='store_true', help='Exercise all real HTTP services without importing Playwright')
    arguments = parser.parse_args()
    run(arguments.base_url, arguments.chromium, http_only=arguments.http_only)
